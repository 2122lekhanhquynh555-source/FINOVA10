import unittest, math
import engine as E, verify, narrative

P = {"income": 25e6, "expenses": 12e6, "investments": 50e6, "assets": 120e6, "debt": 80e6, "debt_payment": 3e6, "years": 10, "liquidity": "medium", "risk": "medium", "debt_rate": 8}


class Independent(unittest.TestCase):
    def test_reference_checks(self):
        for seed in (7, 11, 23):
            good, total, fails = verify.run_checks(seed, 40)
            self.assertEqual(fails, [], fails[:5])
            self.assertEqual(good, total)


class EdgeCases(unittest.TestCase):
    def test_all_zero_and_garbage(self):
        for raw in ({}, None, {"income": 0}, {"income": "abc", "expenses": -5}, {"income": float("nan"), "debt": float("inf")}, {"income": 1e30}):
            a = E.analysis(raw)
            self.assertEqual(len(a["scenarios"]), 5)
            E.whatif(raw, {"income": {"sign": "-", "pct": 100}})
            E.goals_analyze(raw, [{"name": "x", "target": 10, "current": 0, "years": 1}])

    def test_negative_cashflow_draws_down(self):
        a = E.analysis(dict(P, income=8e6))
        base = next(s for s in a["scenarios"] if s["key"] == "baseline")
        self.assertLess(base["net_cash_flow"], 0)
        self.assertGreater(a["scenarios"][0]["shortfall_months"] + base["shortfall_months"], 0)
        self.assertTrue(all(math.isfinite(x) for x in base["net_worth_by_year"]))

    def test_payment_below_interest_flagged(self):
        issues = E.validate_profile(dict(P, debt_payment=100000, debt_rate=24))
        self.assertTrue(any("interest" in i["msg"] for i in issues))
        pl = E.goal_plan("Debt", 80e6, 0, 3, 0, 0, "debt", 24, 100000)
        self.assertTrue(pl["negative_amortization"])

    def test_validation_messages(self):
        msgs = " ".join(i["msg"] for i in E.validate_profile({"income": 0, "debt": 5e6, "debt_payment": 0}))
        self.assertIn("income is missing", msgs.lower()); self.assertIn("no monthly payment", msgs)
        self.assertTrue(any(i["level"] == "error" for i in E.validate_profile({"income": 5e6, "debt": 1, "debt_payment": 9e6})))


class GoalModel(unittest.TestCase):
    def test_no_double_counting(self):
        p = E.clean_profile(P)
        goals = E.default_goals(p)
        self.assertLessEqual(sum(g["current"] for g in goals), p["assets"] + p["investments"])
        bad = [dict(g, current=p["assets"] + p["investments"]) for g in goals[:2]]
        r = E.goals_analyze(P, bad)
        self.assertIn("errors", r); self.assertIn("same money", r["errors"][0])

    def test_capacity_shared_once(self):
        r = E.goals_analyze(dict(P, income=14e6))        # small capacity
        al = r["allocation"]
        self.assertLessEqual(al["total_allocated"], al["capacity"] + 1e-6)
        self.assertAlmostEqual(al["total_allocated"] + al["unallocated"], al["capacity"], 4)

    def test_debt_goal_model(self):
        rd = 0.12 / 12
        pay = E.pmt(100e6, rd, 36)
        pl = E.goal_plan("Debt", 100e6, 0, 3, 0, 0, "debt", 12, pay)
        self.assertEqual(pl["required_monthly"], 0)
        self.assertAlmostEqual(pl["completion"], 1.0, 6)
        self.assertLessEqual(abs(pl["months_at_rate"] - 36), 1)
        pl2 = E.goal_plan("Debt", 100e6, 0, 3, 0, 0, "debt", 12, pay * 0.8)
        self.assertGreater(pl2["required_monthly"], 0); self.assertLess(pl2["completion"], 1)

    def test_required_savings_consistent(self):
        pl = E.goal_plan("G", 500e6, 50e6, 5, 5e6, 0.06)
        self.assertEqual(pl["on_track"], pl["monthly_save"] >= pl["required_monthly"] - 0.5)
        ok = E.goal_plan("G", 500e6, 50e6, 5, pl["required_monthly"], 0.06)
        self.assertAlmostEqual(ok["completion"], 1.0, 6)
        self.assertLessEqual(ok["months_at_rate"], 60)
        self.assertAlmostEqual(ok["milestones"][-1]["balance"], 500e6, 0)

    def test_secondary_goal_and_assumptions(self):
        r = E.goals_analyze(P, None, None, {"income": 40e6, "save_rate": 90, "ret_sign": "-", "ret_pct": 2, "years": 8})
        a = r["assumed"]
        self.assertLessEqual(a["monthly_save"], max(a["net_cash_flow"], 0) + 1e-6)     # cannot save money that does not exist
        self.assertGreaterEqual(a["exceeds_cash_flow"], 0)
        self.assertIsNotNone(r["secondary"]); self.assertNotEqual(r["secondary"]["name"], r["goals"][r["primary"]]["name"])
        one = E.goals_analyze(P, [{"name": "Car", "target": 300e6, "current": 0, "years": 5}], None, {"income": 30e6, "save_rate": 30})
        self.assertTrue(one["secondary"]["suggested"])


def G(name, target, years, current=0, kind="savings", prio="Medium"):
    return {"name": name, "kind": kind, "target": target, "current": current, "years": years, "priority": prio}


def fx(r, name):
    return next(f for f in r["feasibility"] if f["name"] == name)


class Feasibility(unittest.TestCase):
    """The class of a goal must come from real arithmetic, stress tests and shared capacity - never a default 100%."""

    def test_not_every_goal_is_feasible(self):
        r = E.goals_analyze(P, [G("Dream home", 2e9, 5, prio="High"), G("Car", 400e6, 4), G("Trip", 30e6, 3, prio="Low"), G("Safety", 72e6, 2, current=72e6)])
        levels = {f["name"]: f["level"] for f in r["feasibility"]}
        self.assertEqual(levels["Dream home"], "not_feasible"); self.assertEqual(levels["Safety"], "achieved")
        self.assertGreaterEqual(len(set(levels.values())), 3, levels)
        h = fx(r, "Dream home"); self.assertLess(h["completion_base"], 0.5); self.assertTrue(any("needs" in x for x in h["reasons"])); self.assertTrue(h["actions"])
        self.assertTrue(any("lower the target" in a for a in h["actions"]))

    def test_class_boundaries(self):
        # capacity is 13,000,000 / month (no debt); solve targets so required savings are ~20%, ~60%, ~97% and ~140% of capacity
        def level(pct, **kw):
            years, r = 5, E.BASE_RETURN["medium"] / 12
            g, f = E._factor(r, years * 12); target = pct * 13e6 * f
            return E.goals_analyze(dict(P, debt=0, debt_payment=0, **kw), [G("X", target, years)])["feasibility"][0]
        small, mid, tight, over = level(0.20), level(0.60), level(0.97), level(1.40)
        self.assertEqual(small["level"], "highly_feasible", small["scenarios"]["severe"])
        self.assertEqual(mid["level"], "feasible"); self.assertEqual(tight["level"], "at_risk"); self.assertIn(over["level"], ("difficult", "not_feasible"))
        self.assertLess(over["completion_base"], 1); self.assertGreater(over["completion_base"], 0.5)
        for f in (small, mid, tight):
            self.assertAlmostEqual(f["completion_base"], 1.0, 6)

    def test_stress_tests_are_ordered(self):
        import random
        rng = random.Random(5)
        for _ in range(60):
            p = verify.rand_profile(rng); goals = E.default_goals(E.clean_profile(p)) + [G("Big", rng.randint(2, 40) * 1e8, rng.randint(2, 15))]
            r = E.goals_analyze(p, goals)
            for f in r["feasibility"]:
                self.assertLessEqual(f["completion_conservative"], f["completion_base"] + 1e-9, f["name"]); self.assertLessEqual(f["completion_severe"], f["completion_conservative"] + 1e-9)
                self.assertTrue(0 <= f["completion_base"] <= 1)
            for key in ("base", "conservative", "severe"):
                tot = sum(f["scenarios"][key]["allocated"] for f in r["feasibility"]); cap = r["feasibility"][0]["scenarios"][key]["capacity"]
                self.assertLessEqual(tot, cap + 1e-6)

    def test_no_capacity_is_not_feasible_and_explained(self):
        r = E.goals_analyze(dict(P, income=14e6, debt_payment=3e6), [G("Fund", 100e6, 5)]); f = r["feasibility"][0]
        self.assertEqual(f["level"], "not_feasible"); self.assertIn("no money available", " ".join(f["reasons"])); self.assertTrue(any("cash flow" in w for w in r["warnings"]))

    def test_shared_capacity_is_not_double_counted(self):
        years, r_ = 5, E.BASE_RETURN["medium"] / 12; g, f = E._factor(r_, years * 12)
        r = E.goals_analyze(dict(P, debt=0, debt_payment=0), [G("A", 0.7 * 13e6 * f, years, prio="High"), G("B", 0.7 * 13e6 * f, years)], chosen=0)
        a, b = fx(r, "A"), fx(r, "B")
        self.assertEqual(a["level"] in ("at_risk", "feasible", "highly_feasible"), True); self.assertAlmostEqual(a["completion_base"], 1.0, 6)
        self.assertLess(b["completion_base"], 0.5); self.assertIn(b["level"], ("not_feasible", "difficult")); self.assertAlmostEqual(a["allocated_monthly"] + b["allocated_monthly"], 13e6, 3)
        alone = next(x for x in r["ranking"] if x["name"] == "B")["completion"]; self.assertAlmostEqual(alone, 1.0, 6)      # alone it would have looked perfect

    def test_preconditions_cap_the_class(self):
        p = dict(P, assets=10e6, debt=0, debt_payment=0)             # emergency reserve far below 6 months
        years, r_ = 5, E.BASE_RETURN["medium"] / 12; g, f = E._factor(r_, years * 12)
        out = E.goals_analyze(p, [G("Trip", 0.1 * 13e6 * f, years)])["feasibility"][0]
        self.assertEqual(out["level"], "at_risk"); self.assertTrue(any("Emergency reserve" in x for x in out["reasons"]))

    def test_return_dependence_is_disclosed(self):
        out = E.goals_analyze(dict(P, debt=0, debt_payment=0), [G("Fund", 700e6, 8)])["feasibility"][0]
        self.assertIsNotNone(out["completion_no_growth"]); self.assertLess(out["completion_no_growth"], out["completion_base"])
        self.assertTrue(any("0% return" in x for x in out["reasons"]))

    def test_more_income_never_hurts(self):
        rank = {"not_feasible": 0, "difficult": 1, "at_risk": 2, "feasible": 3, "highly_feasible": 4, "achieved": 5}
        goals = [G("Home", 900e6, 7), G("Car", 300e6, 3)]
        a = E.goals_analyze(P, goals)["feasibility"]; b = E.goals_analyze(dict(P, income=35e6), goals)["feasibility"]
        for x, y in zip(a, b): self.assertGreaterEqual(rank[y["level"]], rank[x["level"]], x["name"])

    def test_assumption_and_input_checks(self):
        r = E.goals_analyze(P, [G("Home", 900e6, 50, prio="High")], None, {"income": 90e6, "save_rate": 80, "ret_sign": "+", "ret_pct": 14, "years": 50})
        txt = " ".join(r["warnings"])
        for key in ("high. Returns", "more than double", "unlikely to be sustained", "very long"):
            self.assertIn(key, txt)
        self.assertIn("negative", " ".join(E.goals_analyze(P, [G("A", 5e8, 5)], None, {"ret_sign": "-", "ret_pct": 12})["warnings"]))

    def test_debt_goal_feasibility(self):
        r = E.goals_analyze(dict(P, debt=600e6, debt_payment=2e6, debt_rate=12), [G("Clear debt", 600e6, 3, kind="debt")])
        f = r["feasibility"][0]; self.assertIn(f["level"], ("difficult", "not_feasible", "at_risk")); self.assertTrue(any("extra debt repayment" in x for x in f["reasons"]))

    def test_secondary_goal_has_feasibility(self):
        r = E.goals_analyze(P, [G("Home", 900e6, 7), G("Car", 300e6, 3)], None, {"income": 35e6, "save_rate": 35})
        self.assertIn(r["secondary"]["feasibility"]["level"], E.FEAS_LABEL); self.assertEqual(len(r["feasibility_assumed"]), 2)


class NoOverconfidence(unittest.TestCase):
    """Bad, missing or implausible data and optimistic assumptions must never produce confident-looking percentages."""

    def top(self, a):
        return next(s for s in a["scenarios"] if s["rank"] == 1)

    def test_normal_profile_keeps_its_scores(self):
        a = E.analysis(P); self.assertEqual(self.top(a)["key"], "baseline"); self.assertEqual(self.top(a)["score"], 95); self.assertEqual(self.top(a)["status"], "Most Suitable")
        self.assertEqual(a["confidence"]["notes"], []); self.assertEqual(a["confidence"]["data_quality"]["level"], "high")

    def test_no_income_gives_no_confident_scores(self):
        a = E.analysis(dict(P, income=0))
        self.assertLessEqual(max(s["score"] for s in a["scenarios"]), 10); self.assertTrue(self.top(a)["status"].startswith("Best Available"))
        self.assertEqual(a["confidence"]["data_quality"]["level"], "low"); self.assertTrue(a["confidence"]["notes"])
        r = E.goals_analyze(dict(P, income=0), [G("Home", 500e6, 5), G("Fund", 30e6, 3)])
        for f in r["feasibility"]:
            self.assertEqual(f["level"], "not_feasible"); self.assertEqual(f["completion_base"], 0)
        self.assertLess(max(x["fit"] for x in r["ranking"]), 20)

    def test_missing_expenses_cap_the_class(self):
        goal = [G("Small", 20e6, 5)]
        ok = E.goals_analyze(P, goal)["feasibility"][0]; self.assertEqual(ok["level"], "highly_feasible")
        bad = E.goals_analyze(dict(P, expenses=0), goal); f = bad["feasibility"][0]
        self.assertEqual(bad["data_quality"]["level"], "low"); self.assertEqual(f["level"], "at_risk"); self.assertTrue(any("Data check" in x and "capped at 'At risk'" in x for x in f["reasons"]))
        self.assertLess(bad["ranking"][0]["fit"], E.goals_analyze(P, goal)["ranking"][0]["fit"])
        self.assertTrue(any(i["level"] == "warning" and "expenses" in i["msg"].lower() for i in E.analysis(dict(P, expenses=0))["data_checks"]))

    def test_unusual_inputs_reduce_confidence(self):
        for bad, level in ((dict(P, income=40000), "medium"), (dict(P, expenses=1e6), "medium"), (dict(P, debt_payment=0), "medium"), (dict(P, debt_payment=30e6), "low"),
                           (dict(P, assets=9e12), "medium"), (dict(P, debt=9e11), "medium"), (dict(P, debt_rate=60), "medium")):
            q = E.data_quality(bad); self.assertLessEqual({"low": 0, "medium": 1, "high": 2}[q["level"]], {"low": 0, "medium": 1, "high": 2}[level], bad); self.assertTrue(q["reasons"]); self.assertLess(q["factor"], 1)
        self.assertEqual(E.data_quality(P)["factor"], 1.0)

    def test_negative_cash_flow_is_not_most_suitable(self):
        a = E.analysis(dict(P, income=14e6)); top = self.top(a)
        self.assertLess(max(s["score"] for s in a["scenarios"]), 60); self.assertTrue(top["status"].startswith("Best Available"))
        self.assertTrue(any("negative" in n for n in a["confidence"]["notes"]))

    def test_optimistic_returns_cap_the_class(self):
        goal = [G("Small", 20e6, 5)]
        self.assertEqual(E.goals_analyze(P, goal, None, {"income": 25e6, "save_rate": 40, "ret_sign": "+", "ret_pct": 0, "years": 5})["feasibility_assumed"][0]["level"], "highly_feasible")
        mid = E.goals_analyze(P, goal, None, {"income": 25e6, "save_rate": 40, "ret_sign": "+", "ret_pct": 10, "years": 5})["feasibility_assumed"][0]
        hi = E.goals_analyze(P, goal, None, {"income": 25e6, "save_rate": 40, "ret_sign": "+", "ret_pct": 20, "years": 5})["feasibility_assumed"][0]
        self.assertEqual(mid["level"], "feasible"); self.assertEqual(hi["level"], "at_risk"); self.assertTrue(any("optimistic" in x for x in mid["reasons"] + hi["reasons"]))

    def test_goal_that_only_works_thanks_to_growth_is_not_highly_feasible(self):
        r_, years = E.BASE_RETURN["medium"] / 12, 30
        g, f = E._factor(r_, years * 12)
        out = E.goals_analyze(dict(P, debt=0, debt_payment=0), [G("Far", 0.15 * 13e6 * f, years)])["feasibility"][0]
        self.assertLess(out["completion_no_growth"], 0.5 * out["completion_base"]); self.assertEqual(out["level"], "feasible")
        self.assertTrue(any("depends heavily on investment growth" in x for x in out["reasons"]))

    def test_very_long_timeframe_is_capped(self):
        out = E.goals_analyze(dict(P, debt=0, debt_payment=0), [G("Tiny", 5e6, 45)])["feasibility"][0]
        self.assertNotEqual(out["level"], "highly_feasible"); self.assertTrue(any("too long" in x for x in out["reasons"]))

    def test_shared_capacity_lowers_the_second_goal_not_the_first(self):
        r = E.goals_analyze(P, [G("Dream home", 2e9, 5, prio="High"), G("Car", 400e6, 4)], chosen=1)         # Car is funded first
        alone = next(x for x in r["ranking"] if x["name"] == "Dream home")["completion"]
        self.assertLess(fx(r, "Dream home")["completion_base"], 0.2); self.assertGreater(alone, fx(r, "Dream home")["completion_base"] + 0.1)   # alone it would look better than it is
        self.assertGreater(fx(r, "Car")["completion_base"], 0.95)

    def test_percentages_always_within_bounds_for_random_bad_data(self):
        import random
        rng = random.Random(11)
        for _ in range(80):
            p = {k: rng.choice([0, 0, 1, 1e3, 5e5, 5e6, 3e7, 5e8, 9e12]) for k in ("income", "expenses", "investments", "assets", "debt", "debt_payment")}
            p.update(years=rng.choice([0, 1, 10, 99]), liquidity=rng.choice(E.LEVELS + ("x",)), risk=rng.choice(E.LEVELS), debt_rate=rng.choice([0, 12, 900]))
            a = E.analysis(p)
            for s in a["scenarios"]: self.assertTrue(0 <= s["score"] <= 100)
            r = E.goals_analyze(p, [G("A", rng.choice([1e6, 5e8, 9e12]), rng.choice([1, 10, 50]), current=0, prio=rng.choice(["High", "Low"]))], None, rng.choice([None, {"income": rng.choice([0, 3e7, 9e9]), "save_rate": rng.choice([0, 50, 100, 500]), "ret_pct": rng.choice([0, 50, 200]), "ret_sign": "+"}]))
            if "errors" in r: continue
            for f in r["feasibility"]: self.assertTrue(0 <= f["completion_severe"] <= f["completion_conservative"] + 1e-9 <= f["completion_base"] + 2e-9 <= 1 + 3e-9)
            for x in r["ranking"]: self.assertTrue(0 <= x["fit"] <= 100)


class RequiredData(unittest.TestCase):
    def test_check_required(self):
        self.assertEqual(E.check_required(P), [])
        miss = E.check_required({}); self.assertEqual([x["field"] for x in miss], [k for k, _ in E.REQUIRED]); self.assertTrue(all(x["kind"] == "missing" for x in miss))
        self.assertIn("Enter 0", miss[0]["msg"]); self.assertNotIn("Enter 0", next(x for x in miss if x["field"] == "risk")["msg"])
        self.assertEqual(E.check_required(dict(P, assets=0, debt=0, debt_payment=0)), [])                                    # explicit zero is fine
        self.assertEqual([x["field"] for x in E.check_required(dict(P, assets=None))], ["assets"]); self.assertEqual([x["field"] for x in E.check_required(dict(P, assets="  "))], ["assets"])
        self.assertTrue(E.check_required(dict(P, income="25,000,000 VND")) == [])                                             # typed text with units still parses
        self.assertIn("Missing required information: Assets", E.required_message(E.check_required(dict(P, assets=None))))

    def test_check_goals(self):
        ok = [{"name": "A", "target": 5e8, "years": 5}]
        self.assertEqual(E.check_goals(ok), []); self.assertEqual(E.check_goals(ok + [{"name": "B", "target": 1e6, "years": 1, "current": 0, "priority": "Low"}]), [])
        self.assertEqual(len(E.check_goals(None)), 1); self.assertEqual(len(E.check_goals([])), 1)
        self.assertEqual(len(E.check_goals([{"name": "x"}] * 7)), 1 + 6 * 2)                                                  # too many goals + 2 missing fields each
        self.assertEqual([x["field"] for x in E.check_goals([{"name": "A", "target": 5, "years": 5, "priority": "Urgent"}])], ["priority"])


class WhatIf(unittest.TestCase):
    def test_zero_change_equals_baseline(self):
        w = E.whatif(P, {})
        for r in w["rows"]:
            self.assertTrue(r["variance"] in (None, 0) or abs(r["variance"]) < 1e-6, r)

    def test_each_control_moves_results(self):
        base = E.whatif(P, {})
        d = E.whatif(P, {"debt": {"sign": "+", "pct": 30}})       # faster repayment: less debt left, earlier Debt Freedom
        self.assertLess(next(r for r in d["rows"] if r["label"].startswith("Debt Freedom"))["adjusted"], next(r for r in d["rows"] if r["label"].startswith("Debt Freedom"))["baseline"])
        for key, sign in (("income", "-"), ("expense", "+"), ("unforeseen", "+"), ("savings", "-")):
            w = E.whatif(P, {key: {"sign": sign, "pct": 30}})
            nw = next(r for r in w["rows"] if r["label"].startswith("Net worth"))
            self.assertLess(nw["adjusted"], nw["baseline"], key)
        up = E.whatif(P, {"income": {"sign": "+", "pct": 30}})
        self.assertGreater(next(r for r in up["rows"] if r["label"].startswith("Net worth"))["adjusted"], base["rows"][10]["adjusted"])

    def test_savings_rate_cap(self):
        w = E.whatif(P, {"savings": {"sign": "-", "pct": 100}})
        self.assertAlmostEqual(w["planned"], 0)
        pool = next(r for r in w["rows"] if r["label"].startswith("Investment pool"))
        self.assertLess(pool["adjusted"], pool["baseline"])


class TextAndGrounding(unittest.TestCase):
    def test_grounded_rate(self):
        a = E.analysis(P)
        ok, c = E.grounded_rate("Income is 25,000,000 VND and savings 10,000,000 VND, fake 99,999,999 VND", a)
        self.assertEqual((ok, c), (2, 3))

    def test_narrative_is_grounded(self):
        a = E.analysis(P); n = narrative.analysis(a)
        txt = "\n".join([n["table"], n["chart"], n["guidance"], n["recommendations"]] + list(n["scenarios"].values()))
        ok, c = E.grounded_rate(txt, a)
        self.assertGreaterEqual(ok / c, 0.95)


class ExcelParsing(unittest.TestCase):
    def wb(self, rows, fmt=None, sheets=None):
        from openpyxl import Workbook; import io
        wb = Workbook(); ws = wb.active
        for r in rows: ws.append(r)
        for ref, f in (fmt or {}).items(): ws[ref].number_format = f
        for title, rs in (sheets or {}).items():
            w2 = wb.create_sheet(title)
            for r in rs: w2.append(r)
        b = io.BytesIO(); wb.save(b); return b.getvalue()

    def test_template_roundtrip(self):
        r = E.parse_xlsx(E.template_xlsx())
        self.assertEqual(r["data"]["income"], 25e6); self.assertEqual(r["data"]["risk"], "medium"); self.assertEqual(r["missing"], []); self.assertEqual(r["warnings"], [])

    def test_unusable_files_give_friendly_errors(self):
        for raw in (b"", b"hello", b"PK\x03\x04garbage", b"PK" + b"\x00" * 100, b"x" * (E.MAX_XLSX + 1), b"\xd0\xcf\x11\xe0" + b"0" * 50, self.wb([])):
            with self.assertRaises(ValueError) as cx: E.parse_xlsx(raw)
            self.assertTrue(len(str(cx.exception)) > 10)
        with self.assertRaises(ValueError) as cx: E.parse_xlsx(b"\xd0\xcf\x11\xe0" + b"0" * 50)
        self.assertIn(".xlsx", str(cx.exception))
        r = E.parse_xlsx(self.wb([("random", "stuff"), ("more", 5)])); self.assertEqual(r["data"], {}); self.assertEqual(len(r["unrecognized"]), 2)     # nothing recognised: AI may still map the rows

    def test_vietnamese_labels_units_and_percent_format(self):
        r = E.parse_xlsx(self.wb([("Thu nhập hàng tháng", "25 triệu"), ("Chi tiêu", "12.000.000 đ"), ("Tài sản", "120,000,000 VND"), ("Đầu tư", 50000000), ("Tổng nợ", "80 triệu"),
                                  ("Trả nợ hàng tháng", "3tr"), ("Lãi suất", 0.08), ("Thời hạn", "10 năm"), ("Mức độ rủi ro", "Cao"), ("Thanh khoản", "trung bình")], fmt={"B7": "0%"}))
        self.assertEqual(r["data"], {"income": 25e6, "expenses": 12e6, "assets": 120e6, "investments": 50e6, "debt": 80e6, "debt_payment": 3e6, "debt_rate": 8.0, "years": 10, "risk": "high", "liquidity": "medium"})
        self.assertEqual(r["warnings"], [])

    def test_english_labels_with_brackets_colons_and_case(self):
        r = E.parse_xlsx(self.wb([("MONTHLY INCOME (VND):", 30000000), ("monthly_expenses", 14000000), ("Total Monthly Income note", None), ("Time horizon (years)", 15), ("Risk level", " Low ")]))
        self.assertEqual(r["data"], {"income": 30e6, "expenses": 14e6, "years": 15, "risk": "low"})

    def test_horizontal_table(self):
        r = E.parse_xlsx(self.wb([("Monthly income", "Monthly expenses", "Assets", "Total debt", "Risk level"), (25000000, 12000000, 120000000, 0, "medium")]))
        self.assertEqual(r["data"], {"income": 25e6, "expenses": 12e6, "assets": 120e6, "debt": 0.0, "risk": "medium"})

    def test_blank_value_next_to_example_text_is_not_misread(self):
        r = E.parse_xlsx(self.wb([("Field", "Value", "Example"), ("Monthly income", 25000000, "VND per month"), ("Monthly expenses", None, "VND per month, 0 or more"), ("Assets", "", "VND")]))
        self.assertEqual(r["data"], {"income": 25e6}); self.assertTrue(any("'Monthly expenses' has no value" in w for w in r["warnings"]))
        self.assertFalse(any("not a number" in w for w in r["warnings"])); self.assertIn("expenses", r["missing"])

    def test_title_rows_extra_columns_and_data_on_second_sheet(self):
        r = E.parse_xlsx(self.wb([("Notes",), ("nothing here",)], sheets={"Data": [("My finances 2026",), (), ("Item", "Amount", "Comment"), ("Monthly income", 25000000, "net"), ("Monthly expenses", 12000000, "rent + food")]}))
        self.assertEqual(r["data"], {"income": 25e6, "expenses": 12e6})

    def test_invalid_values_are_reported_not_crashing(self):
        r = E.parse_xlsx(self.wb([("Monthly income", "25,000,000 VND"), ("Monthly expenses", -5), ("Assets", "abc"), ("Risk level", "extreme"), ("Time horizon", 99), ("Monthly income", 1),
                                  ("Total debt", "#REF!"), ("Investments", "=10*2"), ("Debt interest rate", 500)]))
        self.assertEqual(r["data"], {"income": 25e6})
        w = " ".join(r["warnings"])
        for key in ("more than once", "#REF", "negative", "not a number", "low, medium or high", "between 1 and 30", "between 0 and 100", "'Investments' has no value"):
            self.assertIn(key, w, key)

    def test_wrong_sheet_dimension_metadata(self):
        import zipfile, io, re
        raw = self.wb([("Monthly income", 25000000), ("Monthly expenses", 12000000), ("Assets", 5000000)])
        zin = zipfile.ZipFile(io.BytesIO(raw)); out = io.BytesIO(); zout = zipfile.ZipFile(out, "w")
        for n in zin.namelist():
            d = zin.read(n)
            if n.startswith("xl/worksheets/sheet"): d = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', d)
            zout.writestr(n, d)
        zout.close()
        self.assertEqual(E.parse_xlsx(out.getvalue())["data"], {"income": 25e6, "expenses": 12e6, "assets": 5e6})       # files from other tools often lie about their size

    def test_csv_variants(self):
        for raw in ("Monthly income,Monthly expenses\n25000000,12000000\n", "Field;Value\nMonthly income;25.000.000\nMonthly expenses;12.000.000\n",
                    "Monthly income\t25000000\nMonthly expenses\t12000000\n", '\ufeffMonthly income,"25,000,000"\nMonthly expenses,"12,000,000"\n'):
            self.assertEqual(E.parse_xlsx(raw.encode("utf-8"), "data.csv")["data"], {"income": 25e6, "expenses": 12e6}, raw)
        self.assertEqual(E.parse_xlsx("Income;25.000.000\nExpenses;12.000.000\nNote;caf\xe9".encode("cp1252"), "x.csv")["data"], {"income": 25e6, "expenses": 12e6})      # Excel's default Windows encoding
        with self.assertRaises(ValueError): E.parse_xlsx(b"just some text", "notes.pdf")

    def test_unknown_rows_are_offered_for_ai_review(self):
        r = E.parse_xlsx(self.wb([("Monthly income", 25000000), ("Pocket money from granny", 500000)]))
        self.assertEqual(r["unrecognized"], [{"label": "Pocket money from granny", "value": 500000}])

    def test_exports_open_and_are_safe(self):
        from openpyxl import load_workbook; import io
        wb = load_workbook(io.BytesIO(E.report_xlsx(P)))
        self.assertEqual(wb.sheetnames, ["Financial Summary", "Health Assessment", "Scenarios", "Net Worth Projection", "Assumptions"])
        goals = [{"name": "=cmd|' /c calc'!A0", "target": 1e8, "current": 0, "years": 3}]
        wb = load_workbook(io.BytesIO(E.plan_xlsx(P, goals, None, {"income": 30e6})))
        self.assertIn("Funding Allocation", wb.sheetnames)
        self.assertTrue(str(wb["Goals"]["A2"].value).startswith("'"))
        with self.assertRaises(ValueError): E.plan_xlsx(P, [{"name": "x", "target": 0}], None, None)


if __name__ == "__main__":
    unittest.main()
