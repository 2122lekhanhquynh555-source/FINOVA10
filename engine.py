"""FINOVA deterministic calculation engine.

Every financial figure shown by the platform is produced here. The LLM layer only
parses language and explains these outputs; it never calculates them.

Modelling rules (all disclosed in the UI):
  * Cash (the "Assets" field) funds the Emergency Fund goal only.
  * Investments (plus cash above the emergency target) fund the Wealth goal only.
  * Debt is a separate liability amortised with the declared payment and optional interest.
  * Monthly surplus follows a waterfall: emergency cash -> 50% extra debt paydown while debt
    remains -> investments. Each VND of surplus is therefore used exactly once.
"""
import datetime
import io
import math
import re
import time
import unicodedata

BASE_RETURN = {"low": 0.03, "medium": 0.06, "high": 0.09}   # simulated annual yield assumptions
LIQ_MONTHS = {"low": 3, "medium": 6, "high": 9}             # emergency-fund target (months of expenses)
LEVELS = ("low", "medium", "high")
MAX_MONTHS = 600
WEALTH_MONTHS = 36
DEBT_SURPLUS_SHARE = 0.5
MAX_VALUE = 1e13
MONEY_KEYS = ("income", "expenses", "investments", "assets", "debt", "debt_payment")


# ----------------------------------------------------------------- formatting
def vnd(x):
    return f"{x:,.0f} VND"


def pc(x, d=1):
    if x is None:
        return "n/a"
    if d == 0 and 0.995 <= x < 1:                      # never show 99.6% as a full 100%
        d = 1
    return f"{x * 100:.{d}f}%"


def dur(m):
    if m is None:
        return "not reached within 50 years"
    y, mo = divmod(int(round(m)), 12)
    if y and mo:
        return f"{y} yr {mo} mo"
    if y:
        return f"{y} yr"
    return f"{mo} mo" if mo else "already achieved"


# ------------------------------------------------------------------- profile
def _num(v):
    try:
        if isinstance(v, bool):
            return 0.0
        v = float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        v = 0.0
    return v if math.isfinite(v) else 0.0


def _pos(v):
    return min(max(_num(v), 0.0), MAX_VALUE)


def clean_profile(d):
    d = d or {}
    p = {k: _pos(d.get(k)) for k in MONEY_KEYS}
    p["debt_rate"] = min(max(_num(d.get("debt_rate")), 0.0), 100.0)      # % per year
    p["years"] = int(min(max(_num(d.get("years", 10)) or 10, 1), 30))
    for k in ("liquidity", "risk"):
        v = str(d.get(k, "medium")).lower()
        p[k] = v if v in LEVELS else "medium"
    return p


REQUIRED = [("income", "Monthly income"), ("expenses", "Monthly expenses"), ("investments", "Investments"), ("assets", "Assets"),
            ("debt", "Total current debt"), ("debt_payment", "Monthly debt payment"), ("years", "Time horizon (years)"),
            ("liquidity", "Liquidity need"), ("risk", "Risk level")]


def check_required(raw):
    """Strict validation of what the user typed. Returns a list of {field, label, kind, msg}; empty list = complete and valid.
    A blank box is NOT the same as 0: a missing value is an error, an explicit 0 is accepted."""
    raw = raw or {}
    errs = []

    def bad(k, label, kind, msg):
        errs.append({"field": k, "label": label, "kind": kind, "msg": msg})

    for k, label in REQUIRED:
        v = raw.get(k)
        if v is None or (isinstance(v, str) and not v.strip()):
            bad(k, label, "missing", f"{label} is required." + ("" if k in ("liquidity", "risk") else " Enter 0 if it does not apply."))
            continue
        if k in ("liquidity", "risk"):
            if str(v).strip().lower() not in LEVELS:
                bad(k, label, "invalid", f"{label} must be low, medium or high.")
            continue
        x = _clean_number(v)
        if x is None:
            bad(k, label, "invalid", f"{label} must be a number (you typed {str(v)[:20]!r}).")
        elif x < 0:
            bad(k, label, "invalid", f"{label} cannot be negative.")
        elif x > MAX_VALUE:
            bad(k, label, "invalid", f"{label} is unrealistically large.")
        elif k == "income" and x <= 0:
            bad(k, label, "invalid", "Monthly income must be greater than 0.")
        elif k == "years" and (not 1 <= x <= 30 or x != int(x)):
            bad(k, label, "invalid", "Time horizon must be a whole number of years between 1 and 30.")
    v = raw.get("debt_rate")
    if v is not None and not (isinstance(v, str) and not v.strip()):
        x = _clean_number(v)
        if x is None or not 0 <= x <= 100:
            bad("debt_rate", "Debt interest rate", "invalid", "Debt interest rate must be a number from 0 to 100 (% per year).")
    if not any(e["field"] in ("income", "debt_payment") for e in errs):
        inc, dp = _clean_number(raw.get("income")) or 0, _clean_number(raw.get("debt_payment")) or 0
        if dp > inc > 0:
            bad("debt_payment", "Monthly debt payment", "invalid", "The monthly debt payment cannot be larger than monthly income.")
    return errs


def required_message(errs):
    miss = [e["label"] for e in errs if e["kind"] == "missing"]
    parts = []
    if miss:
        parts.append("Missing required information: " + ", ".join(miss) + ". Enter 0 where something does not apply.")
    parts += [e["msg"] for e in errs if e["kind"] == "invalid"]
    return " ".join(parts)


def validate_profile(d):
    """Python-side validation used by every input path (form, text parsing, Excel)."""
    raw = d or {}
    p = clean_profile(raw)
    issues = []

    def add(level, field, msg):
        issues.append({"level": level, "field": field, "msg": msg})

    for k in MONEY_KEYS:
        if k in raw and _num(raw[k]) < 0:
            add("warning", k, f"{k.replace('_', ' ')} was negative and has been set to 0.")
        if k in raw and _num(raw[k]) > MAX_VALUE:
            add("error", k, f"{k.replace('_', ' ')} is unrealistically large.")
    if not p["income"]:
        add("error", "income", "Monthly income is missing or 0; ratios and projections cannot be computed.")
    if not p["expenses"]:
        add("warning", "expenses", "Monthly expenses are 0 - please confirm this is correct.")
    if p["debt"] > 0 and not p["debt_payment"]:
        add("warning", "debt_payment", "You have debt but no monthly payment; the debt would never be repaid.")
    if p["debt"] == 0 and p["debt_payment"] > 0:
        add("warning", "debt", "A monthly debt payment was entered but total debt is 0 (conflicting data).")
    if p["income"] and p["debt_payment"] > p["income"]:
        add("error", "debt_payment", "Monthly debt payment is larger than monthly income.")
    elif p["income"] and p["expenses"] + p["debt_payment"] > p["income"]:
        add("warning", "income", "Expenses plus debt payments exceed income, so cash flow is negative.")
    if p["debt"] > 0 and p["debt_payment"] > 0 and p["debt_payment"] <= p["debt"] * p["debt_rate"] / 1200:
        add("warning", "debt_payment", "The monthly payment does not cover the monthly interest; the debt will not shrink.")
    for k in ("years", "liquidity", "risk"):
        if k not in raw:
            add("info", k, f"{k} was not provided; the default ({p[k]}) is used.")
    return issues


def metrics(p):
    inc, exp = p["income"], p["expenses"]
    net = inc - exp - p["debt_payment"]
    total_assets = p["assets"] + p["investments"]
    return {
        "income": inc, "expenses": exp, "debt_payment": p["debt_payment"],
        "net_cash_flow": net,
        "expense_ratio": exp / inc if inc else None,
        "dti": p["debt_payment"] / inc if inc else None,
        "savings_capacity": max(net, 0),
        "savings_rate": max(net, 0) / inc if inc else 0,
        "total_assets": total_assets,
        "liabilities": p["debt"],
        "net_worth": total_assets - p["debt"],
        "emergency_months": p["assets"] / exp if exp else None,
        "target_months": LIQ_MONTHS[p["liquidity"]],
    }


def health(p, m):
    inc = m["income"]
    items = []
    if inc <= 0:
        cf = ("bad", "No income", "No income was entered, so ongoing living expenses and debt service cannot be covered.")
    elif m["net_cash_flow"] >= 0.10 * inc:
        cf = ("good", "Positive", f"Income covers living expenses and debt service with {vnd(m['net_cash_flow'])} left each month ({pc(m['net_cash_flow'] / inc)} of income).")
    elif m["net_cash_flow"] >= 0:
        cf = ("warn", "Balanced (thin margin)", f"Income only just covers expenses and debt service; the monthly margin is {vnd(m['net_cash_flow'])}.")
    else:
        cf = ("bad", "Negative", f"Expenses plus debt service exceed income by {vnd(-m['net_cash_flow'])} per month.")
    items.append(("Cash Flow Analysis",) + cf)
    d = m["dti"]
    if not p["debt_payment"]:
        db = ("good", "No debt service", "No monthly debt payments were entered.")
    elif d is None:
        db = ("bad", "Unmeasurable", "Debt payments exist but no income was entered.")
    elif d < 0.20:
        db = ("good", "Low", f"Debt payments take {pc(d)} of income, below the commonly used 20% comfort level.")
    elif d < 0.36:
        db = ("warn", "Moderate", f"Debt payments take {pc(d)} of income, inside the 20%-36% watch zone.")
    elif d < 0.50:
        db = ("bad", "High", f"Debt payments take {pc(d)} of income, above the 36% guideline.")
    else:
        db = ("bad", "Severe", f"Debt payments take {pc(d)} of income, which leaves very little room for other needs.")
    items.append(("Debt Burden",) + db)
    em, tm = m["emergency_months"], m["target_months"]
    if em is None:
        lq = ("warn", "Unmeasurable", "No living expenses were entered, so reserve coverage cannot be measured.")
    elif em >= tm:
        lq = ("good", "Strong", f"Assets cover {em:.1f} months of expenses, meeting your {tm}-month target.")
    elif em >= tm / 2:
        lq = ("warn", "Partial", f"Assets cover {em:.1f} months of expenses against a {tm}-month target.")
    else:
        lq = ("bad", "Thin", f"Assets cover only {em:.1f} months of expenses against a {tm}-month target.")
    items.append(("Liquidity",) + lq)
    sr = m["savings_rate"]
    if sr >= 0.20:
        sv = ("good", "Strong", f"You can set aside about {pc(sr)} of income ({vnd(m['savings_capacity'])} per month).")
    elif sr >= 0.10:
        sv = ("good", "Healthy", f"You can set aside about {pc(sr)} of income ({vnd(m['savings_capacity'])} per month).")
    elif sr >= 0.05:
        sv = ("warn", "Modest", f"Savings capacity is {pc(sr)} of income ({vnd(m['savings_capacity'])} per month).")
    else:
        sv = ("bad", "Low", f"Savings capacity is only {pc(sr)} of income ({vnd(m['savings_capacity'])} per month).")
    items.append(("Savings Capacity",) + sv)

    strengths, weak, actions = [], [], []
    for name, lvl, status, txt in items:
        (strengths if lvl == "good" else weak if lvl == "bad" else []).append(f"{name}: {status}")
    if m["net_cash_flow"] < 0:
        actions.append("Close the monthly cash-flow gap first (reduce expenses, raise income, or restructure debt).")
    if d is not None and d >= 0.36:
        actions.append("Lower the debt burden: prioritise high-cost debt and avoid new borrowing.")
    if em is not None and em < tm:
        actions.append(f"Build the emergency reserve toward {tm} months of expenses ({vnd(tm * m['expenses'])}).")
    if sr < 0.10 and m["net_cash_flow"] >= 0:
        actions.append("Raise the savings rate toward 10%-20% of income.")
    if not actions:
        actions.append("Maintain current habits and move on to selecting savings or investment goals.")
    return {"items": [{"pillar": a, "level": b, "status": c, "text": t} for a, b, c, t in items],
            "strengths": strengths, "vulnerabilities": weak, "actions": actions}


def data_quality(profile):
    """How far can the inputs be trusted? Implausible or missing data must never produce confident-looking percentages.
    Returns {"level": high|medium|low, "factor": 0.3-1.0, "reasons": [...]}."""
    raw = profile or {}
    p = clean_profile(raw)
    inc, exp = p["income"], p["expenses"]
    rank = {"high": 2, "medium": 1, "low": 0}
    level, factor, reasons = "high", 1.0, []

    def hit(lv, f, msg):
        nonlocal level, factor
        if rank[lv] < rank[level]:
            level = lv
        factor *= f
        reasons.append(msg)

    if inc <= 0:
        hit("low", 0.4, "No monthly income was entered, so capacity and goal results cannot be trusted.")
    else:
        if inc < 500_000:
            hit("medium", 0.8, f"Monthly income ({vnd(inc)}) is unusually low - amounts must be entered in VND, not in millions.")
        if exp <= 0:
            hit("low", 0.6, "Monthly expenses are 0, so your savings capacity is probably overstated.")
        elif exp / inc < 0.10:
            hit("medium", 0.85, f"Expenses are only {pc(exp / inc, 0)} of income, which is unusually low; savings capacity may be overstated.")
        if p["debt_payment"] > inc:
            hit("low", 0.5, "The monthly debt payment is larger than monthly income.")
    if p["debt"] > 0 and p["debt_payment"] <= 0:
        hit("medium", 0.8, "You have debt but no monthly debt payment, so the debt plan cannot be trusted.")
    elif p["debt"] > 0 and p["debt_payment"] <= p["debt"] * p["debt_rate"] / 1200:
        hit("medium", 0.85, "The debt payment does not cover the monthly interest.")
    if p["debt"] == 0 and p["debt_payment"] > 0:
        hit("medium", 0.9, "A debt payment was entered but total debt is 0 (conflicting data).")
    if inc > 0 and p["assets"] + p["investments"] > 600 * inc:
        hit("medium", 0.85, "Assets plus investments exceed 50 years of income, which is unusual; please check the amounts.")
    if inc > 0 and p["debt"] > 240 * inc:
        hit("medium", 0.85, "Debt exceeds 20 years of income; please check the amounts.")
    if p["debt_rate"] > 40:
        hit("medium", 0.9, f"The debt interest rate ({p['debt_rate']:g}%) is unusually high.")
    return {"level": level, "factor": max(factor, 0.3), "reasons": reasons}


def health_factor(p, m):
    """Scales suitability and fit scores by real financial health (negative cash flow, heavy debt, thin reserve)."""
    f, why = 1.0, []
    if p["income"] <= 0:
        return 0.2, ["No income: nothing can be funded."]
    if m["net_cash_flow"] < 0:
        f *= 0.6
        why.append(f"Monthly cash flow is negative ({vnd(m['net_cash_flow'])}).")
    elif m["savings_rate"] < 0.05:
        f *= 0.8
        why.append(f"Savings capacity is only {pc(m['savings_rate'])} of income.")
    if m["dti"] is not None and m["dti"] >= 0.5:
        f *= 0.7
        why.append(f"Debt payments take {pc(m['dti'])} of income.")
    elif m["dti"] is not None and m["dti"] >= 0.36:
        f *= 0.85
        why.append(f"Debt payments take {pc(m['dti'])} of income (above 36%).")
    if m["emergency_months"] is not None and m["emergency_months"] < m["target_months"] / 2:
        f *= 0.9
        why.append(f"The emergency reserve covers only {m['emergency_months']:.1f} of the {m['target_months']} target months.")
    return f, why


# --------------------------------------------------------------- simulation
def _run(p, inc, exp, ret, extra=0.0, debt_pay=None, save_cap=None, months=MAX_MONTHS):
    """Month-by-month waterfall simulation (see module docstring).

    inc/exp: monthly income and living expenses; ret: annual yield on investments;
    extra: unforeseen monthly expense; debt_pay: scheduled payment override;
    save_cap: optional monthly cap on money saved (the rest is spent)."""
    r = max(ret, -0.9) / 12
    rd = p["debt_rate"] / 1200
    sched = p["debt_payment"] if debt_pay is None else debt_pay
    em_t = LIQ_MONTHS[p["liquidity"]] * p["expenses"]
    wt = WEALTH_MONTHS * p["expenses"]
    cash, pool, debt = p["assets"], p["investments"], p["debt"]
    reach = {"g1": None, "g2": None, "g3": None}
    short = 0

    def chk(m):
        if reach["g1"] is None and cash >= em_t - 1e-9:
            reach["g1"] = m
        if reach["g2"] is None and debt <= 0.5:
            reach["g2"] = m
        if reach["g3"] is None and pool >= wt - 1e-9:
            reach["g3"] = m

    ys = {"assets": [cash + pool], "nw": [cash + pool - debt], "pool": [pool], "cash": [cash], "debt": [debt]}
    cf0 = None
    chk(0)
    for m in range(1, months + 1):
        interest = debt * rd
        pay = min(sched, debt + interest)
        debt = debt + interest - pay
        if debt < 0.5:
            debt = 0.0
        cf = inc - exp - extra - pay
        if cf0 is None:
            cf0 = cf
        if pool > 0:
            pool *= 1 + r
        if cf >= 0:
            s = cf if save_cap is None else min(cf, save_cap)
            if cash < em_t:
                add = min(s, em_t - cash)
                cash += add
                s -= add
            if debt > 0 and s > 0:
                x = min(s * DEBT_SURPLUS_SHARE, debt)
                debt -= x
                s -= x
            pool += s
        else:
            cash += cf
            if cash < 0:
                pool += cash
                cash = 0.0
        if pool < 0:
            short += 1
        chk(m)
        if m % 12 == 0:
            ys["assets"].append(cash + pool)
            ys["nw"].append(cash + pool - debt)
            ys["pool"].append(pool)
            ys["cash"].append(cash)
            ys["debt"].append(debt)
    return {"cf0": cf0, "reach": reach, "shortfall_months": short, **ys}


def simulate(p, sc):
    ret = BASE_RETURN[p["risk"]] + sc["ret"]
    inc = p["income"] * (1 + sc["income"])
    exp = p["expenses"] * (1 + sc["expense"])
    run = _run(p, inc, exp, ret)
    return {"key": sc["key"], "name": sc["name"], "short": sc["short"], "desc": sc["desc"],
            "assumptions": {"income_change": sc["income"], "expense_change": sc["expense"], "annual_return": max(ret, -0.9)},
            "monthly_income": inc, "monthly_expenses": exp, "net_cash_flow": run["cf0"],
            "savings_capacity": max(run["cf0"], 0), "assets_by_year": run["assets"], "net_worth_by_year": run["nw"],
            "pool_by_year": run["pool"], "debt_by_year": run["debt"], "reach": run["reach"],
            "shortfall_months": run["shortfall_months"]}


SCENARIOS = [
    {"key": "severe", "name": "Severely Conservative Scenario", "short": "Severely Conservative",
     "income": -0.25, "expense": 0.15, "ret": -0.04,
     "desc": "Sharp income drop (-25%), major expense surge (+15%) and investment yield 4 points below the base assumption."},
    {"key": "conservative", "name": "Conservative Scenario", "short": "Conservative",
     "income": -0.10, "expense": 0.05, "ret": -0.02,
     "desc": "Mild income decline (-10%), slight expense increase (+5%) and yield 2 points lower."},
    {"key": "baseline", "name": "Baseline Scenario", "short": "Baseline",
     "income": 0.0, "expense": 0.0, "ret": 0.0,
     "desc": "Your own income, expenses and debt payments with the base yield for your risk level."},
    {"key": "optimistic", "name": "Optimistic Scenario", "short": "Optimistic",
     "income": 0.10, "expense": 0.0, "ret": 0.01,
     "desc": "Income 10% higher than baseline, other declared assumptions unchanged, yield 1 point higher."},
    {"key": "high", "name": "Highly Optimistic Scenario", "short": "Highly Optimistic",
     "income": 0.25, "expense": 0.0, "ret": 0.02,
     "desc": "Significant growth: income +25% and yield 2 points higher, expenses unchanged."},
]
BASE_SCORE = {"severe": 40, "conservative": 85, "baseline": 95, "optimistic": 78, "high": 55}
RISK_ADJ = {"low": {"severe": 5, "conservative": 5, "baseline": 0, "optimistic": -6, "high": -10},
            "medium": {"severe": 0, "conservative": 0, "baseline": 0, "optimistic": 0, "high": 0},
            "high": {"severe": -5, "conservative": -5, "baseline": 0, "optimistic": 4, "high": 8}}


def _label(rank, score):
    if rank == 1:
        return "Most Suitable" if score >= 60 else "Best Available (Low Confidence)"
    return "Suitable" if score >= 80 else "Moderately Suitable" if score >= 60 else \
        "Use with Caution" if score >= 45 else "Risk Contingency"


def goal_targets(p):
    return {"g1": LIQ_MONTHS[p["liquidity"]] * p["expenses"], "g2": p["debt"], "g3": WEALTH_MONTHS * p["expenses"]}


def analysis(profile):
    p = clean_profile(profile)
    m = metrics(p)
    sims = [simulate(p, sc) for sc in SCENARIOS]
    base = next(s for s in sims if s["key"] == "baseline")
    n = p["years"]
    tg = goal_targets(p)
    dq = data_quality(profile)
    hf, hf_why = health_factor(p, m)
    for s in sims:
        sc = BASE_SCORE[s["key"]] + RISK_ADJ[p["risk"]][s["key"]]
        if m["net_cash_flow"] < 0:
            sc += {"baseline": -20, "optimistic": -15, "high": -20, "conservative": 3, "severe": 10}[s["key"]]
        if m["savings_rate"] >= 0.20:
            sc += {"optimistic": 4, "high": 3}.get(s["key"], 0)
        s["raw_score"] = int(max(0, min(100, sc)))
        s["score"] = int(max(0, min(100, round(s["raw_score"] * hf * dq["factor"]))))      # weak health or doubtful data lowers every score
        s["final_net_worth"] = s["net_worth_by_year"][n]
        s["final_assets"] = s["assets_by_year"][n]
        s["final_pool"] = s["pool_by_year"][n]
        s["goal_delay"] = {}
        for g in ("g1", "g2", "g3"):
            a, b = s["reach"][g], base["reach"][g]
            s["goal_delay"][g] = None if a is None or b is None else a - b
        s["wealth_gap"] = s["final_pool"] - tg["g3"]
    ranked = sorted(sims, key=lambda s: -s["score"])
    for i, s in enumerate(ranked, 1):
        s["rank"] = i
        s["status"] = _label(i, s["score"])
        s["title"] = f"{i}. {s['name']} ({s['status']} — {s['score']}%)"
    return {"profile": p, "metrics": m, "health": health(p, m), "scenarios": sims,
            "ranking": [s["key"] for s in ranked], "goals": default_goals(p), "goal_targets": tg, "data_checks": [i for i in validate_profile(profile) if i["level"] != "info"],
            "confidence": {"data_quality": dq, "health_factor": hf, "health_reasons": hf_why, "scale": hf * dq["factor"],
                           "notes": ([f"Suitability scores are scaled down to {pc(hf * dq['factor'], 0)} of their raw value because of the points below."] + hf_why + dq["reasons"]) if hf * dq["factor"] < 0.999 else []},
            "years": list(range(0, n + 1)),
            "assumption_notes": [
                "Monthly surplus follows a waterfall: (1) refill the cash emergency reserve, (2) 50% to extra debt repayment while debt remains, (3) the rest is invested at the scenario's simulated annual yield. Each VND is used once.",
                "Cash (Assets) is the emergency reserve; investments are the invested pool. They are never counted for two goals.",
                f"Debt is amortised monthly with the declared payment and the interest rate you entered ({p['debt_rate']:.1f}% per year; 0% if left blank).",
                "Deficits are drawn from cash first, then from the investment pool.",
                f"Base annual yield by risk level: low {pc(BASE_RETURN['low'], 0)}, medium {pc(BASE_RETURN['medium'], 0)}, high {pc(BASE_RETURN['high'], 0)}; simulated assumptions, not guarantees.",
                f"Reference goals: Emergency Fund (cash = {LIQ_MONTHS[p['liquidity']]} months of expenses), Debt Freedom, Wealth Milestone (investment pool = {WEALTH_MONTHS} months of expenses)."]}


# ------------------------------------------------------------------ what-if
def _adj(adj, key):
    a = (adj or {}).get(key) or {}
    s = -1 if str(a.get("sign", "+")) in ("-", "−", "-1", "minus") else 1
    return max(s * min(max(_num(a.get("pct")), 0), 1000) / 100, -1.0)


def whatif(profile, adj):
    p = clean_profile(profile)
    m = metrics(p)
    n = p["years"]
    fi, fe, fu, fs, fd = (_adj(adj, k) for k in ("income", "expense", "unforeseen", "savings", "debt"))
    inc2 = p["income"] * (1 + fi)
    exp2 = p["expenses"] * (1 + fe)
    unf2 = max(fu, 0) * p["expenses"]            # '+' = emergency shock; '-' = no emergency spending
    debt2 = p["debt_payment"] * (1 + fd)
    net2 = inc2 - exp2 - unf2 - debt2
    planned = m["savings_capacity"] * (1 + fs)
    cap = None if fs == 0 else max(planned, 0)
    ret = BASE_RETURN[p["risk"]]
    b, a = _run(p, p["income"], p["expenses"], ret), _run(p, inc2, exp2, ret, extra=unf2, debt_pay=debt2, save_cap=cap)
    tm = m["target_months"]
    rows = [
        ("Monthly income", p["income"], inc2, "VND"),
        ("Living expenses", p["expenses"], exp2, "VND"),
        ("Unforeseen expenses", 0, unf2, "VND"),
        ("Monthly debt service", p["debt_payment"], debt2, "VND"),
        ("Net monthly cash flow", m["net_cash_flow"], net2, "VND"),
        ("Savings capacity", m["savings_capacity"], max(net2, 0), "VND"),
        ("Planned savings (after savings-rate change)", m["savings_capacity"], planned, "VND"),
        ("Savings feasibility gap (capacity - plan)", 0, max(net2, 0) - planned, "VND"),
        ("Liquidity surplus / (deficit) vs target reserve", p["assets"] - tm * p["expenses"], p["assets"] - tm * (exp2 + unf2), "VND"),
        ("Reserve coverage (months)", m["emergency_months"] or 0, p["assets"] / (exp2 + unf2) if exp2 + unf2 else 0, "months"),
        (f"Net worth after {n} years", b["nw"][n], a["nw"][n], "VND"),
        (f"Investment pool after {n} years", b["pool"][n], a["pool"][n], "VND"),
        (f"Remaining debt after {n} years", b["debt"][n], a["debt"][n], "VND"),
        ("Emergency Fund reached", b["reach"]["g1"], a["reach"]["g1"], "dur"),
        ("Debt Freedom reached", b["reach"]["g2"], a["reach"]["g2"], "dur"),
        ("Wealth Milestone reached", b["reach"]["g3"], a["reach"]["g3"], "dur"),
    ]
    out = []
    for lab, x, y, u in rows:
        out.append({"label": lab, "baseline": x, "adjusted": y, "unit": u,
                    "variance": (None if x is None or y is None else y - x)})
    warns = []
    if net2 < 0:
        warns.append("Net cash flow is negative after the adjustments: savings and assets would be drawn down.")
    if a["shortfall_months"]:
        warns.append(f"The investment pool would be depleted for {a['shortfall_months']} months (funding shortfall).")
    return {"adjustments": {"income": fi, "expense": fe, "unforeseen": fu, "savings": fs, "debt": fd},
            "rows": out,
            "chart": {"labels": ["Net cash flow", "Savings capacity", "Planned savings"],
                      "baseline": [m["net_cash_flow"], m["savings_capacity"], m["savings_capacity"]],
                      "adjusted": [net2, max(net2, 0), planned]},
            "series": {"labels": list(range(0, n + 1)), "baseline": b["nw"][:n + 1], "adjusted": a["nw"][:n + 1]},
            "net": net2, "base_net": m["net_cash_flow"], "liquidity": p["assets"] - tm * (exp2 + unf2),
            "planned": planned, "warnings": warns}


# -------------------------------------------------------------------- goals
def _factor(r, n):
    g = (1 + r) ** n
    return g, (n if r == 0 else (g - 1) / r)


def pmt(pv, r, n):
    """Level payment that amortises pv over n periods at periodic rate r."""
    if pv <= 0:
        return 0.0
    return pv / n if r == 0 else pv * r / (1 - (1 + r) ** -n)


def nper(pv, r, pay):
    """Months to repay pv with payment pay at periodic rate r (None if never)."""
    if pv <= 0.5:
        return 0
    if pay <= 0:
        return None
    if r == 0:
        return math.ceil(pv / pay - 1e-9)
    if pay <= pv * r:
        return None
    return math.ceil(-math.log(1 - pv * r / pay) / math.log(1 + r) - 1e-9)


def goal_plan(name, target, current, years, monthly_save, annual_ret, kind="savings", debt_rate=0.0, sched_payment=0.0):
    if kind == "debt":
        return _debt_plan(name, target, current, years, monthly_save, debt_rate, sched_payment)
    n = max(int(round(years * 12)), 1)
    r = annual_ret / 12
    g, fvf = _factor(r, n)
    projected = current * g + monthly_save * fvf
    need = max(target - current * g, 0)
    req = need / fvf if fvf > 0 else 0.0
    months_at = 0 if current >= target else None
    bal = current
    if months_at is None:
        for k in range(1, 1201):
            bal = bal * (1 + r) + monthly_save
            if bal >= target - 1e-9:
                months_at = k
                break
    completion = 1.0 if target <= 0 else min(projected / target, 1.0)
    series, rseries, b1, b2 = [current], [current], current, current
    for k in range(1, n + 1):
        b1 = b1 * (1 + r) + monthly_save
        b2 = b2 * (1 + r) + req
        if k % 12 == 0:
            series.append(b1)
            rseries.append(b2)
    miles = []
    for q in (0.25, 0.5, 0.75, 1.0):
        k = max(int(round(n * q)), 1)
        gk, fk = _factor(r, k)
        miles.append({"pct": q, "month": k, "balance": current * gk + req * fk, "projected": current * gk + monthly_save * fk})
    return {"kind": "savings", "name": name, "target": target, "current": current, "years": years, "monthly_save": monthly_save,
            "annual_return": annual_ret, "extra_needed": max(target - current, 0), "required_monthly": req,
            "months_at_rate": months_at, "gap": monthly_save - req, "projected": projected, "completion": completion,
            "shortfall": max(target - projected, 0), "surplus": max(projected - target, 0),
            "on_track": monthly_save >= req - 0.5,
            "options": {"extend_years": None if months_at is None else math.ceil(months_at / 12),
                        "reduced_target": min(projected, target), "required_monthly": req},
            "milestones": miles, "series": series, "required_series": rseries}


def _debt_plan(name, target, current, years, extra, debt_rate, sched):
    """Debt goal: `extra` is the additional monthly money (savings capacity) applied on top of the
    scheduled payment. `current` is money already set aside against the principal."""
    D = max(target - current, 0.0)
    rd = debt_rate / 1200
    n = max(int(round(years * 12)), 1)
    req_total = pmt(D, rd, n)
    req_extra = max(req_total - sched, 0.0)
    pay = sched + extra
    bal = D
    months_at = 0 if D <= 0.5 else None
    rem_n = 0.0 if D <= 0.5 else None
    series = [target - D]
    for k in range(1, 1201):
        if bal > 0:
            bal = bal * (1 + rd) - pay
            if bal <= 0.5:
                bal = 0.0
                if months_at is None:
                    months_at = k
        if k == n:
            rem_n = bal
        if k % 12 == 0 and k <= n:
            series.append(target - bal)
        if k >= n and (months_at is not None or pay <= bal * rd):
            break
    if rem_n is None:
        rem_n = bal
    projected = max(target - rem_n, 0.0)          # a debt that keeps growing has repaid nothing
    miles = []
    for q in (0.25, 0.5, 0.75, 1.0):
        kk = max(int(round(n * q)), 1)
        g_, f_ = _factor(rd, kk)
        b_ = max(D * g_ - req_total * f_, 0.0)
        miles.append({"pct": q, "month": kk, "balance": target - b_, "projected": None})
    completion = 1.0 if target <= 0 else min(projected / target, 1.0)
    return {"kind": "debt", "name": name, "target": target, "current": current, "years": years, "monthly_save": extra,
            "annual_return": 0.0, "debt_rate": debt_rate, "scheduled_payment": sched, "extra_needed": D,
            "required_monthly": req_extra, "required_total_payment": req_total, "months_at_rate": months_at,
            "gap": extra - req_extra, "projected": projected, "completion": completion,
            "shortfall": max(target - projected, 0), "surplus": 0.0, "on_track": extra >= req_extra - 0.5,
            "negative_amortization": pay + 1e-9 < D * rd,
            "options": {"extend_years": None if months_at is None else math.ceil(months_at / 12),
                        "reduced_target": min(projected, target), "required_monthly": req_extra},
            "milestones": miles, "series": series, "required_series": []}


PRIORITY = {"high": 100, "medium": 70, "low": 40}


def default_goals(p):
    m = metrics(p)
    tm = LIQ_MONTHS[p["liquidity"]]
    em = tm * p["expenses"]
    em_cur = min(p["assets"], em)
    spare = p["assets"] - em_cur
    g = [{"name": "Emergency Fund", "kind": "savings", "target": round(em), "current": int(math.floor(em_cur)),
          "years": 2, "priority": "High", "constraints": f"Cover {tm} months of living expenses (cash)"}]
    if p["debt"] > 0:
        mo = nper(p["debt"], p["debt_rate"] / 1200, p["debt_payment"])
        yrs = min(max(math.ceil(mo / 12), 1), 30) if mo else 5
        g.append({"name": "Debt Freedom", "kind": "debt", "target": round(p["debt"]), "current": 0, "years": yrs,
                  "priority": "Medium", "constraints": "Clear all current liabilities"})
    else:
        g.append({"name": "Home / Education Fund", "kind": "savings", "target": round(12 * p["income"]), "current": 0,
                  "years": max(p["years"] // 2, 3), "priority": "Medium", "constraints": "New savings equal to 12 months of income"})
    g.append({"name": "Wealth Milestone", "kind": "savings", "target": round(WEALTH_MONTHS * p["expenses"]),
              "current": int(math.floor(p["investments"] + spare)), "years": p["years"], "priority": "Medium",
              "constraints": f"Invested funds equal to {WEALTH_MONTHS} months of expenses"})
    g = [x for x in g if x["target"] > 0]          # a 0 target (e.g. expenses not entered) is not a goal
    return g or [{"name": "My savings goal", "kind": "savings", "target": 0, "current": 0, "years": 5, "priority": "Medium", "constraints": "Enter a target amount"}]


def check_goals(raw):
    """Strict check of the goal form. A goal needs a name, a target amount above 0 and a timeframe; blank is an error.
    Returns [{goal, field, msg}] with `goal` = 0-based index of the goal card (None for a general message)."""
    if not isinstance(raw, list) or not raw:
        return [{"goal": None, "field": None, "msg": "Add at least one goal."}]
    errs = []
    if len(raw) > 6:
        errs.append({"goal": None, "field": None, "msg": "You can analyse up to 6 goals at a time."})
    blank = lambda v: v is None or (isinstance(v, str) and not v.strip())
    for i, g in enumerate(raw[:6]):
        n = i + 1
        if not isinstance(g, dict):
            errs.append({"goal": i, "field": "name", "msg": f"Goal {n} is not valid."})
            continue
        if not str(g.get("name") or "").strip():
            errs.append({"goal": i, "field": "name", "msg": f"Goal {n}: please enter a goal name."})
        if blank(g.get("target")):
            errs.append({"goal": i, "field": "target", "msg": f"Goal {n}: the target amount is required."})
        else:
            x = _clean_number(g.get("target"))
            if x is None or x <= 0 or x > MAX_VALUE:
                errs.append({"goal": i, "field": "target", "msg": f"Goal {n}: the target amount must be a number greater than 0."})
        if blank(g.get("years")):
            errs.append({"goal": i, "field": "years", "msg": f"Goal {n}: the target timeframe (years) is required."})
        else:
            y = _clean_number(g.get("years"))
            if y is None or not 1 <= y <= 50:
                errs.append({"goal": i, "field": "years", "msg": f"Goal {n}: the timeframe must be between 1 and 50 years."})
        if not blank(g.get("current")):
            c = _clean_number(g.get("current"))
            if c is None or c < 0 or c > MAX_VALUE:
                errs.append({"goal": i, "field": "current", "msg": f"Goal {n}: existing funds must be 0 or more."})
        if g.get("priority") not in (None, "", "High", "Medium", "Low"):
            errs.append({"goal": i, "field": "priority", "msg": f"Goal {n}: priority must be High, Medium or Low."})
    return errs


def clean_goals(goals, keep_blank=False):
    out = []
    for g in (goals or [])[:6]:
        if not isinstance(g, dict):
            continue
        name = str(g.get("name", "")).strip()[:80]
        if not name and not keep_blank:
            continue
        kind = "debt" if str(g.get("kind", "")).lower() == "debt" else "savings"
        out.append({"name": name, "kind": kind, "target": _pos(g.get("target")), "current": _pos(g.get("current")),
                    "years": min(max(_num(g.get("years", 5)) or 5, 1), 50),
                    "priority": g.get("priority") if g.get("priority") in ("High", "Medium", "Low") else "Medium",
                    "constraints": str(g.get("constraints", ""))[:200]})
    return out


def clean_assumptions(a):
    """Sanitise the hypothetical-assumption inputs of the goal what-if tool (None when empty/invalid)."""
    if not isinstance(a, dict) or not a:
        return None
    return {"years": min(max(_num(a.get("years", 5)) or 5, 1), 50), "income": _pos(a.get("income")),
            "save_rate": min(max(_num(a.get("save_rate")), 0), 100),
            "ret_sign": "-" if str(a.get("ret_sign", "+")) in ("-", "−") else "+", "ret_pct": min(max(_num(a.get("ret_pct")), 0), 100)}


def clean_adjustments(adj):
    """Sanitise the five stress-test inputs: {key: {sign, pct}}."""
    out = {}
    for k in ("income", "expense", "unforeseen", "savings", "debt"):
        a = (adj or {}).get(k) or {}
        out[k] = {"sign": "-" if str(a.get("sign", "+")) in ("-", "−", "-1", "minus") else "+", "pct": min(max(_num(a.get("pct")), 0), 1000)}
    return out


def validate_goals(p, goals):
    errors, warnings = [], []
    if not goals:
        errors.append("Add at least one goal with a name and a target amount.")
    for g in goals:
        if g["target"] <= 0:
            errors.append(f"'{g['name']}': the target amount must be above 0.")
    names = [g["name"].lower() for g in goals]
    if len(set(names)) != len(names):
        warnings.append("Two goals share the same name; consider renaming one.")
    avail = p["assets"] + p["investments"]
    used = sum(g["current"] for g in goals)
    if used > avail + 1:
        errors.append(f"Existing funds assigned to goals ({vnd(used)}) exceed your total assets plus investments ({vnd(avail)}). "
                      "The same money cannot fund two goals - reduce the existing funds of one or more goals.")
    for g in goals:
        if g["kind"] == "debt" and p["debt"] > 0 and g["target"] > p["debt"] * 1.0001:
            warnings.append(f"'{g['name']}': the target is larger than your recorded total debt ({vnd(p['debt'])}).")
        if g["kind"] == "debt" and p["debt"] == 0:
            warnings.append(f"'{g['name']}': you recorded no debt in Financial Information.")
    if sum(1 for g in goals if g["kind"] == "debt") > 1:
        warnings.append("Several debt goals are not combined: each one is calculated against the same debt payment.")
    return errors, warnings


def _fit(p, m, goal, plan):
    notes, pen = [], 0
    is_em = "emergency" in goal["name"].lower()
    is_debt = goal["kind"] == "debt"
    if m["emergency_months"] is not None and m["emergency_months"] < m["target_months"] and not is_em:
        pen += 10
        notes.append(f"Emergency reserve is below the {m['target_months']}-month target (precondition not met).")
    if m["dti"] is not None and m["dti"] >= 0.36 and not is_debt:
        pen += 10
        notes.append("Debt-to-income is above 36%, so debt reduction should be handled in parallel.")
    if goal["years"] <= 2 and p["risk"] == "high":
        notes.append("Short timeframe: avoid high-volatility assets for this goal.")
    if plan.get("negative_amortization"):
        pen += 15
        notes.append("The payment does not cover the interest, so this debt would grow.")
    if goal["target"] > 0 and goal["current"] >= goal["target"]:
        pen += 60
        notes.append("This goal is already fully funded, so it needs no further saving.")
    pr = PRIORITY.get(str(goal.get("priority", "Medium")).lower(), 70)
    return max(0, min(100, round(0.7 * plan["completion"] * 100 + 0.3 * pr - pen))), notes


def _make_plan(p, g, years, save, ret):
    return goal_plan(g["name"], g["target"], g["current"], years, save, ret, g["kind"], p["debt_rate"], p["debt_payment"])


def _rank(p, m, goals, monthly_save, ret, years_override=None, idx_override=None, scale=1.0):
    out = []
    for i, g in enumerate(goals):
        y = years_override if (years_override and i == idx_override) else g["years"]
        pl = _make_plan(p, g, y, monthly_save, ret)
        fit, notes = _fit(p, m, g, pl)
        fit = int(round(fit * scale))                    # doubtful data or weak financial health lowers the fit score
        out.append({"idx": i, "name": g["name"], "kind": g["kind"], "fit": fit, "completion": pl["completion"],
                    "required_monthly": pl["required_monthly"],
                    "share_of_capacity": (pl["required_monthly"] / monthly_save) if monthly_save else None,
                    "priority": g.get("priority", "Medium"), "years": y, "target": g["target"], "notes": notes, "plan": pl})
    return sorted(out, key=lambda x: (-x["fit"], -x["completion"], x["idx"]))


def allocate(ranked, capacity):
    """Share one monthly capacity across goals in ranking order so no VND is counted twice."""
    rem = capacity
    rows, tot_req = [], 0.0
    for x in ranked:
        req = x["plan"]["required_monthly"]
        got = min(req, rem)
        rem -= got
        tot_req += req
        rows.append({"idx": x["idx"], "name": x["name"], "required": req, "allocated": got,
                     "funded": 1.0 if req <= 0 else got / req})
    return {"rows": rows, "capacity": capacity, "total_required": tot_req, "total_allocated": capacity - rem,
            "unallocated": rem, "shortfall": max(tot_req - capacity, 0.0)}


STRESS = [("base", "Baseline", 0.0, 0.0, 0.0), ("conservative", "Conservative stress", -0.10, 0.05, -0.02), ("severe", "Severe stress", -0.25, 0.15, -0.04)]
FEAS_LABEL = {"achieved": "Already achieved", "highly_feasible": "Highly feasible", "feasible": "Feasible", "at_risk": "At risk",
              "difficult": "Difficult", "not_feasible": "Not feasible"}


def stress_info(income, expenses, base_capacity, base_ret):
    out = []
    for key, name, di, de, dr in STRESS:
        out.append({"key": key, "name": name, "income_change": di, "expense_change": de, "return_change": dr,
                    "capacity": max(base_capacity + income * di - expenses * de, 0.0), "annual_return": max(base_ret + dr, -0.9)})
    return out


FEAS_ORDER = ["not_feasible", "difficult", "at_risk", "feasible", "highly_feasible", "achieved"]


def _cap_level(level, cap):
    return level if FEAS_ORDER.index(level) <= FEAS_ORDER.index(cap) else cap


def feasibility(p, m, goals, order, years_map, income, base_capacity, base_ret, alone, dq=None):
    """Realistic feasibility of every goal.

    One monthly capacity is shared by all goals in `order` (so no VND is counted twice). Each goal is evaluated under three
    conditions - baseline, conservative stress (income -10%, expenses +5%, yield -2 pts) and severe stress (-25%, +15%, -4 pts) -
    because requirements rise and capacity shrinks when conditions worsen. The class is derived from those results:
      not_feasible  baseline completion < 50% (or no capacity)      difficult   baseline completion < 90%
      at_risk       baseline 90-100%, or conservative stress < 90%  feasible    survives conservative stress
      highly_feasible  also survives severe stress (>= 90%)         achieved    existing funds already cover the target"""
    stress = stress_info(income, p["expenses"], base_capacity, base_ret)
    table = {}
    for st in stress:
        rem, cap, ret, per = st["capacity"], st["capacity"], st["annual_return"], {}
        for i in order:
            g, y = goals[i], years_map.get(i, goals[i]["years"])
            req = _make_plan(p, g, y, 0.0, ret)["required_monthly"]
            got = min(req, rem)
            rem -= got
            pl = _make_plan(p, g, y, got, ret)
            per[i] = {"capacity": cap, "required": req, "allocated": got, "funded": 1.0 if req <= 0 else got / req, "completion": pl["completion"],
                      "months": pl["months_at_rate"], "projected": pl["projected"], "shortfall": pl["shortfall"]}
        table[st["key"]] = per
    out = []
    for pos, i in enumerate(order):
        g, y = goals[i], years_map.get(i, goals[i]["years"])
        b, c, sv = table["base"][i], table["conservative"][i], table["severe"][i]
        n_months = max(int(round(y * 12)), 1)
        pl_b = _make_plan(p, g, y, b["allocated"], base_ret)
        gap = max(b["required"] - b["allocated"], 0.0)
        ng = _make_plan(p, g, y, b["allocated"], 0.0)["completion"] if g["kind"] == "savings" else None
        need = g["target"] - g["current"] > 0.5
        reasons, actions = [], []
        if not need:
            level = "achieved"
            reasons.append(f"Existing funds of {vnd(g['current'])} already cover the target of {vnd(g['target'])}.")
        elif base_capacity <= 0:
            level = "not_feasible"
            reasons.append(f"Your monthly cash flow after expenses and debt payments is {vnd(m['net_cash_flow'])}, so there is no money available to save toward this goal.")
        elif b["completion"] < 0.5 or pl_b["months_at_rate"] is None:
            level = "not_feasible"
        elif b["completion"] < 0.9:
            level = "difficult"
        elif b["completion"] < 0.9999 or c["completion"] < 0.9:
            level = "at_risk"
        elif sv["completion"] < 0.9:
            level = "feasible"
        else:
            level = "highly_feasible"
        if need and base_capacity > 0:
            what = "extra debt repayment above your scheduled payment" if g["kind"] == "debt" else "savings"
            reasons.append(f"Reaching {vnd(g['target'])} in {y:g} years needs {vnd(b['required'])} of {what} per month (existing funds: {vnd(g['current'])}). "
                           f"Your shared monthly capacity is {vnd(b['capacity'])}; {vnd(b['allocated'])} of it is available to this goal after the goals ranked before it ({pc(b['funded'], 0)} of the requirement), "
                           f"giving an estimated completion of {pc(b['completion'], 0)}.")
            if income and b["required"] > 0.3 * income:
                reasons.append(f"The requirement equals {pc(b['required'] / income, 0)} of your monthly income, which is a heavy burden to sustain.")
            if b["required"] > b["capacity"] > 0:
                reasons.append(f"On its own this goal needs more than your entire capacity ({vnd(b['required'])} vs {vnd(b['capacity'])}).")
            if level in ("highly_feasible", "feasible", "at_risk") or c["completion"] < b["completion"] - 1e-9:
                reasons.append(f"If income falls 10%, expenses rise 5% and yields are 2 points lower, capacity drops to {vnd(c['capacity'])} and completion changes from "
                               f"{pc(b['completion'], 0)} to {pc(c['completion'], 0)}; under severe stress (income -25%, expenses +15%, yields -4 points) it is {pc(sv['completion'], 0)}.")
            if ng is not None and base_ret > 0 and b["allocated"] > 0:
                growth = max(pl_b["projected"] - g["current"] - b["allocated"] * n_months, 0.0)
                reasons.append(f"About {vnd(growth)} of the projected {vnd(pl_b['projected'])} depends on the assumed {pc(base_ret)} yearly return; with a 0% return the same savings would reach {pc(ng, 0)} of the target.")
            if pl_b["months_at_rate"] is None:
                reasons.append("At the amount available to this goal, the target is not reached within 50 years.")
            for note in _fit(p, m, g, pl_b)[1]:
                reasons.append(note)
            # a good-looking result is never allowed to rest on doubtful data or on optimistic assumptions
            if level in ("feasible", "highly_feasible"):
                caps = []
                if dq and dq["level"] == "low":
                    caps.append(("at_risk", "Data check: " + " ".join(dq["reasons"]) + " Results are capped at 'At risk' until the data is corrected."))
                elif dq and dq["level"] == "medium":
                    caps.append(("feasible", "Data check: " + " ".join(dq["reasons"]) + " Results are capped at 'Feasible' until the data is confirmed."))
                if base_ret > 0.20:
                    caps.append(("at_risk", f"The assumed return of {pc(base_ret)} per year is extremely optimistic, so the result is capped at 'At risk'."))
                elif base_ret > 0.12:
                    caps.append(("feasible", f"The assumed return of {pc(base_ret)} per year is optimistic, so the result is capped at 'Feasible'."))
                if ng is not None and base_ret > 0 and level == "highly_feasible" and ng < 0.5 * b["completion"]:
                    caps.append(("feasible", f"Reaching the goal depends heavily on investment growth (only {pc(ng, 0)} would be reached with a 0% return), so it is not rated 'Highly feasible'."))
                if y > 40:
                    caps.append(("feasible", f"A timeframe of {y:g} years is too long to rate any goal 'Highly feasible'."))
                for cap, why_ in caps:
                    new = _cap_level(level, cap)
                    if new != level:
                        level = new
                    reasons.append(why_)
            if level in ("highly_feasible", "feasible") and any("precondition" in x or "above 36%" in x for x in reasons):
                level = "at_risk"
            if gap > 0.5:
                if pl_b["months_at_rate"] is not None:
                    actions.append(f"Extend the timeframe to about {math.ceil(pl_b['months_at_rate'] / 12)} years, or")
                actions.append(f"lower the target to about {vnd(pl_b['projected'])}, or")
                actions.append(f"free up about {vnd(gap)} more per month (raise income or cut expenses).")
            if level == "at_risk" and gap <= 0.5:
                actions.append("Keep a buffer: the plan works today but loses ground when conditions worsen. Build the emergency reserve and avoid new debt.")
            if g["kind"] == "savings" and g["years"] <= 2 and p["risk"] == "high":
                actions.append("Use low-volatility assets for such a short timeframe.")
        out.append({"idx": i, "name": g["name"], "kind": g["kind"], "order": pos + 1, "level": level, "label": FEAS_LABEL[level],
                    "required_monthly": b["required"], "allocated_monthly": b["allocated"], "funded": b["funded"],
                    "completion_base": b["completion"], "completion_conservative": c["completion"], "completion_severe": sv["completion"],
                    "completion_no_growth": ng, "scenarios": {k: table[k][i] for k in table}, "reasons": reasons, "actions": actions})
    return sorted(out, key=lambda x: x["idx"]), stress


def assumption_checks(p, m, goals, ret, assumed=None, base_income=None):
    """Reasonableness checks on the inputs and assumptions; returned as plain-language warnings."""
    w = []
    if p["income"] <= 0:
        w.append("No monthly income was entered, so no goal can be assessed realistically.")
    elif m["net_cash_flow"] <= 0:
        w.append(f"Your monthly cash flow is {vnd(m['net_cash_flow'])}: expenses and debt payments use all of your income, so nothing is available for goals until that improves.")
    elif m["savings_rate"] < 0.05:
        w.append(f"Your savings capacity is only {pc(m['savings_rate'])} of income, so most goals will be hard to reach.")
    if p["debt"] > 0 and p["debt_payment"] > 0 and p["debt_payment"] <= p["debt"] * p["debt_rate"] / 1200:
        w.append("Your debt payment does not cover the monthly interest, so the debt would grow instead of shrinking.")
    if ret > 0.12:
        w.append(f"The assumed annual return of {pc(ret)} is high. Returns this large are not guaranteed and rarely hold for decades, so treat these results as optimistic.")
    if ret < 0:
        w.append(f"The assumed annual return is negative ({pc(ret)}); results show a capital loss scenario.")
    if assumed:
        base = p["income"] if base_income is None else base_income
        if base and assumed["monthly_income"] > 2 * base:
            w.append(f"The hypothetical income ({vnd(assumed['monthly_income'])}) is more than double your current income ({vnd(base)}); this is a hypothesis, not a promise.")
        if assumed["savings_rate"] > 0.5:
            w.append(f"Saving {pc(assumed['savings_rate'], 0)} of income is unlikely to be sustained for many years.")
        if assumed["exceeds_cash_flow"] > 0:
            w.append(f"The requested savings exceed the assumed available cash flow by {vnd(assumed['exceeds_cash_flow'])}; the result uses only the money that exists.")
    for g in goals:
        if g["years"] > 40:
            w.append(f"'{g['name']}': a timeframe of {g['years']:g} years is very long, so inflation and life changes matter more than the calculation shows.")
        if p["income"] and g["target"] > 600 * p["income"]:
            w.append(f"'{g['name']}': the target equals more than 50 years of your current income.")
        if g["target"] > 0 and g["current"] >= g["target"]:
            w.append(f"'{g['name']}': existing funds already cover the target.")
    return w


def _slim(r):
    return {k: v for k, v in r.items() if k != "plan"}


def goals_analyze(profile, goals=None, chosen=None, assumptions=None):
    p = clean_profile(profile)
    m = metrics(p)
    goals = clean_goals(goals) or default_goals(p)
    errors, warnings = validate_goals(p, goals)
    if errors:
        return {"errors": errors, "warnings": warnings}
    base_ret = BASE_RETURN[p["risk"]]
    base_save = m["savings_capacity"]
    dq = data_quality(profile)
    hf, _hf_why = health_factor(p, m)
    scale = dq["factor"] * hf
    ranked = _rank(p, m, goals, base_save, base_ret, scale=scale)
    ai_pick = ranked[0]["idx"]
    primary = chosen if isinstance(chosen, int) and 0 <= chosen < len(goals) else ai_pick
    pg = goals[primary]
    plan_base = _make_plan(p, pg, pg["years"], base_save, base_ret)
    order = [primary] + [x["idx"] for x in ranked if x["idx"] != primary]      # the goal the client works on is funded first
    by_idx = {x["idx"]: x for x in ranked}
    alloc = allocate([by_idx[i] for i in order], base_save)
    feas, stress = feasibility(p, m, goals, order, {}, p["income"], base_save, base_ret, {i: x["plan"] for i, x in by_idx.items()}, dq)
    warnings = warnings + assumption_checks(p, m, goals, base_ret)
    res = {"profile": p, "metrics": m, "goals": goals, "warnings": warnings, "feasibility": feas, "stress": stress, "data_quality": dq, "scale": scale,
           "feasibility_assumed": None,
           "ranking": [_slim(r) | {"rank": i + 1} for i, r in enumerate(ranked)],
           "ai_pick": ai_pick, "primary": primary, "plan_baseline": plan_base, "allocation": alloc,
           "baseline_capacity": {"monthly_income": p["income"], "net_cash_flow": m["net_cash_flow"], "monthly_save": base_save,
                                 "annual_return": base_ret, "savings_rate": m["savings_rate"]},
           "assumed": None, "secondary": None, "plan_assumed": None, "allocation_assumed": None}
    a = assumptions or None
    if a:
        inc2 = _pos(a.get("income")) or p["income"]
        rate = min(max(_num(a.get("save_rate", m["savings_rate"] * 100)), 0), 100) / 100
        sign = -1 if str(a.get("ret_sign", "+")) in ("-", "−") else 1
        ret2 = max(base_ret + sign * min(max(_num(a.get("ret_pct")), 0), 100) / 100, -0.9)
        yrs2 = min(max(_num(a.get("years", pg["years"])) or pg["years"], 1), 50)
        cf2 = inc2 - p["expenses"] - p["debt_payment"]
        req_save = rate * inc2
        save2 = min(req_save, max(cf2, 0))
        res["assumed"] = {"monthly_income": inc2, "net_cash_flow": cf2, "requested_save": req_save, "monthly_save": save2,
                          "annual_return": ret2, "savings_rate": rate, "years": yrs2, "exceeds_cash_flow": max(req_save - save2, 0)}
        res["plan_assumed"] = _make_plan(p, pg, yrs2, save2, ret2)
        p2 = dict(p, income=inc2)
        m2 = metrics(p2)
        r2_all = _rank(p2, m2, goals, save2, ret2, yrs2, primary, scale)
        order2 = [primary] + [x["idx"] for x in r2_all if x["idx"] != primary]
        by2 = {x["idx"]: x for x in r2_all}
        res["allocation_assumed"] = allocate([by2[i] for i in order2], save2)
        res["feasibility_assumed"], res["stress_assumed"] = feasibility(p2, m2, goals, order2, {primary: yrs2}, inc2, save2, ret2, {i: x["plan"] for i, x in by2.items()}, dq)
        res["warnings"] = res["warnings"] + [w for w in assumption_checks(p2, m2, goals, ret2, res["assumed"], p["income"]) if w not in res["warnings"]]
        others = [g for i, g in enumerate(goals) if i != primary]
        suggested = False
        if not others:
            others = [{"name": "Additional Investment Fund", "kind": "savings", "target": round(12 * inc2), "current": 0,
                       "years": max(int(yrs2), 5), "priority": "Medium", "constraints": "AI-suggested optional goal"}]
            suggested = True
        r2 = _rank(p2, m2, others, save2, ret2, scale=scale)[0]
        sg = others[r2["idx"]]
        res["secondary"] = {"name": sg["name"], "kind": sg["kind"], "suggested": suggested, "fit": r2["fit"], "notes": r2["notes"],
                            "plan": r2["plan"], "years": sg["years"], "plan_baseline": _make_plan(p, sg, sg["years"], base_save, base_ret)}
        if suggested:
            res["secondary"]["feasibility"] = feasibility(p2, m2, [sg], [0], {}, inc2, save2, ret2, {0: r2["plan"]}, dq)[0][0]
        else:
            orig = [i for i in range(len(goals)) if i != primary][r2["idx"]]
            res["secondary"]["feasibility"] = next(f for f in res["feasibility_assumed"] if f["idx"] == orig)
    ylen = max(len(plan_base["series"]), len(res["plan_assumed"]["series"]) if res["plan_assumed"] else 0)
    pad = lambda s: s + [None] * (ylen - len(s))
    res["chart"] = {"labels": list(range(ylen)), "baseline": pad(plan_base["series"]),
                    "required": pad(plan_base["required_series"]) if plan_base["required_series"] else None,
                    "assumed": pad(res["plan_assumed"]["series"]) if res["plan_assumed"] else None,
                    "target": [pg["target"]] * ylen}
    return res


# ------------------------------------------------------------------- Excel
def _sheet(wb, title, rows, widths=(44, 26, 26, 26, 26)):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, float):
                c.number_format = "#,##0.00" if abs(c.value) < 10 else "#,##0"
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="41734E")
    return ws


DISCLAIMER = ("FINOVA is a financial planning support platform for educational and informational purposes only. "
              "Results are mathematical estimates based on user-provided data and stated assumptions, not guarantees, "
              "audits or investment advice.")


def _safe(v):
    """Neutralise spreadsheet formula injection in exported text."""
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


def report_xlsx(profile):
    from openpyxl import Workbook
    a = analysis(profile)
    p, m = a["profile"], a["metrics"]
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Financial Summary", [
        ("Item", "Value", "Unit"),
        ("Monthly income", p["income"], "VND/month"), ("Monthly expenses", p["expenses"], "VND/month"),
        ("Monthly debt payment", p["debt_payment"], "VND/month"), ("Investments", p["investments"], "VND"),
        ("Assets", p["assets"], "VND"), ("Total debt", p["debt"], "VND"), ("Debt interest rate", p["debt_rate"], "% per year"),
        ("Time horizon", p["years"], "years"), ("Liquidity need", p["liquidity"], ""), ("Risk level", p["risk"], ""),
        ("Net monthly cash flow", m["net_cash_flow"], "VND/month"),
        ("Expense-to-income ratio", m["expense_ratio"] if m["expense_ratio"] is not None else "n/a", "ratio"),
        ("Debt-to-income (DTI)", m["dti"] if m["dti"] is not None else "n/a", "ratio"),
        ("Monthly savings capacity", m["savings_capacity"], "VND/month"), ("Total assets", m["total_assets"], "VND"),
        ("Total current liabilities", m["liabilities"], "VND"), ("Estimated net worth", m["net_worth"], "VND"),
        (), ("Disclaimer", DISCLAIMER)], widths=(46, 70, 16))
    _sheet(wb, "Health Assessment", [("Pillar", "Status", "Detail")] + [(i["pillar"], i["status"], i["text"]) for i in a["health"]["items"]]
           + [(), ("Actions",)] + [(x,) for x in a["health"]["actions"]], widths=(30, 24, 90))
    rows = [("Scenario", "Suitability %", "Net cash flow / month", "Savings capacity / month", "Annual yield", "Final net worth", "Assumptions")]
    for s in sorted(a["scenarios"], key=lambda s: s["rank"]):
        rows.append((s["title"], s["score"], s["net_cash_flow"], s["savings_capacity"], s["assumptions"]["annual_return"], s["final_net_worth"], s["desc"]))
    _sheet(wb, "Scenarios", rows, widths=(52, 14, 22, 24, 14, 22, 80))
    rows = [("Year",) + tuple(s["short"] for s in a["scenarios"])]
    for i, y in enumerate(a["years"]):
        rows.append((y,) + tuple(s["net_worth_by_year"][i] for s in a["scenarios"]))
    _sheet(wb, "Net Worth Projection", rows, widths=(10, 24, 24, 24, 24, 24))
    _sheet(wb, "Assumptions", [("Assumption",)] + [(x,) for x in a["assumption_notes"]], widths=(120,))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def plan_xlsx(profile, goals, chosen, assumptions):
    from openpyxl import Workbook
    r = goals_analyze(profile, goals, chosen, assumptions)
    if "errors" in r:
        raise ValueError(" ".join(r["errors"]))
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Goals", [("Goal", "Type", "Target (VND)", "Existing funds (VND)", "Years", "Priority", "Constraints")] +
           [(_safe(g["name"]), g["kind"], g["target"], g["current"], g["years"], g["priority"], _safe(g["constraints"])) for g in r["goals"]],
           widths=(30, 10, 20, 22, 10, 12, 50))
    _sheet(wb, "AI Ranking", [("Rank", "Goal", "Fit score", "Completion", "Required monthly (VND)")] +
           [(x["rank"], _safe(x["name"]), x["fit"], x["completion"], x["required_monthly"]) for x in r["ranking"]], widths=(8, 34, 12, 14, 26))
    al = r["allocation"]
    _sheet(wb, "Feasibility", [("Goal", "Class", "Completion (baseline, shared capacity)", "Conservative stress", "Severe stress", "Required / month", "Allocated / month", "Why", "What would help")] +
           [(_safe(f["name"]), f["label"], f["completion_base"], f["completion_conservative"], f["completion_severe"], f["required_monthly"], f["allocated_monthly"],
             " ".join(f["reasons"]), " ".join(f["actions"])) for f in r["feasibility"]], widths=(26, 18, 22, 18, 16, 20, 20, 90, 70))
    _sheet(wb, "Funding Allocation", [("Goal", "Required / month", "Allocated / month", "Funded %")] +
           [(_safe(x["name"]), x["required"], x["allocated"], x["funded"]) for x in al["rows"]] +
           [(), ("Capacity", al["capacity"]), ("Unallocated", al["unallocated"]), ("Shortfall", al["shortfall"])], widths=(34, 22, 22, 14))

    def plan_rows(title, pl):
        return [(title,), ("Goal", _safe(pl["name"])), ("Type", pl["kind"]), ("Target", pl["target"]), ("Existing funds", pl["current"]), ("Years", pl["years"]),
                ("Monthly amount applied", pl["monthly_save"]), ("Annual return (simulated)", pl["annual_return"]),
                ("Additional total required", pl["extra_needed"]), ("Required monthly", pl["required_monthly"]),
                ("Time to goal at this rate", dur(pl["months_at_rate"])), ("Gap (applied - required)", pl["gap"]),
                ("Projected value at deadline", pl["projected"]), ("Completion (if all capacity goes to this goal)", pl["completion"]), ()]
    rows = plan_rows("Plan on current baseline", r["plan_baseline"])
    if r["plan_assumed"]:
        rows += plan_rows("Plan on future assumptions (hypothetical)", r["plan_assumed"])
    if r["secondary"]:
        rows += plan_rows("Secondary goal (hypothetical)", r["secondary"]["plan"])
    rows += [("Disclaimer", DISCLAIMER)]
    _sheet(wb, "Plans", rows, widths=(40, 40))
    c = r["chart"]
    rows = [("Year", "Baseline plan", "Required path", "Assumed plan", "Target")] + [
        (y, c["baseline"][i], (c["required"] or [None] * 99)[i], (c["assumed"] or [None] * 99)[i], c["target"][i]) for i, y in enumerate(c["labels"])]
    _sheet(wb, "Projection", rows, widths=(10, 22, 22, 22, 22))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


ALIASES = {
    "income": ["monthly income", "income", "total monthly income", "net income", "salary", "monthly salary", "thu nhập", "thu nhập hàng tháng", "thu nhập tháng", "tổng thu nhập", "lương", "lương tháng"],
    "expenses": ["monthly expenses", "expenses", "monthly spending", "living expenses", "monthly expense", "chi tiêu", "chi tiêu hàng tháng", "chi phí", "chi phí hàng tháng", "chi phí sinh hoạt"],
    "investments": ["investments", "investment", "invested assets", "đầu tư", "khoản đầu tư", "tổng đầu tư"],
    "assets": ["assets", "savings", "cash", "cash and savings", "tài sản", "tiết kiệm", "tiền mặt", "tiền gửi"],
    "years": ["time horizon (years)", "time horizon", "years", "horizon", "planning horizon", "thời hạn", "thời gian", "kỳ hạn", "số năm"],
    "liquidity": ["liquidity need", "liquidity", "liquidity needs", "nhu cầu thanh khoản", "thanh khoản"],
    "risk": ["risk level", "risk", "risk tolerance", "risk appetite", "mức độ rủi ro", "rủi ro"],
    "debt": ["total current debt", "total debt", "debt", "liabilities", "loans", "tổng nợ", "nợ", "khoản vay", "dư nợ"],
    "debt_payment": ["monthly debt payment", "debt payment", "monthly debt service", "debt service", "loan payment", "nợ phải trả hàng tháng", "trả nợ hàng tháng", "tiền trả nợ"],
    "debt_rate": ["debt interest rate", "average debt interest rate", "interest rate", "loan interest rate", "lãi suất nợ", "lãi suất"],
}


def _fold(s):
    """Lower-case, accent-free, punctuation-free form used to compare spreadsheet labels."""
    s = str(s).strip().lower().replace("đ", "d")
    s = "".join(ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn")
    s = re.sub(r"\(.*?\)|\[.*?\]", " ", s)
    s = re.sub(r"[_:*/\\\-]+", " ", s)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9% ]", " ", s)).strip()


FIELD_MAP = {_fold(a): k for k, v in ALIASES.items() for a in v}
CATEGORY_WORDS = {"low": "low", "thấp": "low", "thap": "low", "medium": "medium", "moderate": "medium", "mid": "medium", "trung bình": "medium", "trung binh": "medium",
                  "high": "high", "cao": "high"}
MAX_XLSX = 4 * 1024 * 1024
HEADER_WORDS = {"field", "item", "label", "name", "value", "values", "example", "examples", "notes", "note", "type", "unit", "units", "stt", "muc", "gia tri", "ghi chu"}


def field_for_label(label):
    """Map a spreadsheet label to a FINOVA field (exact alias first, then keyword rules). None if unknown."""
    l = _fold(label)
    if not l or l in HEADER_WORDS or len(l) > 70:
        return None
    if l in FIELD_MAP:
        return FIELD_MAP[l]
    words = set(l.split())
    if "horizon" in l or "ky han" in l or "thoi han" in l or "thoi gian" in l or "so nam" in l or l in ("year", "years", "nam"):
        return "years"
    if "liquid" in l or "thanh khoan" in l:
        return "liquidity"
    if "risk" in l or "rui ro" in l:
        return "risk"
    if any(w in l for w in ("debt", "loan", "liabilit", "vay", "no phai tra")) or l in ("no", "tong no", "khoan no", "du no"):
        if "rate" in l or "interest" in l or "lai" in words:
            return "debt_rate"
        if "payment" in l or "service" in l or "monthly" in l or "tra no" in l or "hang thang" in l or "installment" in l:
            return "debt_payment"
        return "debt"
    if any(w in l for w in ("income", "salary", "wage", "thu nhap", "luong")):
        return "income"
    if any(w in l for w in ("expens", "spend", "chi tieu", "chi phi")):
        return "expenses"
    if "invest" in l or "dau tu" in l:
        return "investments"
    if any(w in l for w in ("asset", "saving", "cash", "tai san", "tiet kiem", "tien mat")):
        return "assets"
    return None


_UNIT_X = {"billion": 1e9, "bn": 1e9, "b": 1e9, "ty": 1e9, "tỷ": 1e9, "tỉ": 1e9, "million": 1e6, "mil": 1e6, "m": 1e6, "trieu": 1e6, "triệu": 1e6, "tr": 1e6,
           "thousand": 1e3, "k": 1e3, "nghin": 1e3, "nghìn": 1e3, "ngan": 1e3, "ngàn": 1e3}
_NUM_RE = re.compile(r"([-+]?\s*\d[\d.,\s]*)\s*(billion|bn|tỷ|tỉ|ty|million|mil|triệu|trieu|tr|thousand|nghìn|nghin|ngàn|ngan|k|m|b)?\s*"
                     r"(vnd|vnđ|₫|đ|d|%|/ ?month|/ ?thang|/ ?tháng|per month|a month|mỗi tháng|hàng tháng|/ ?year|/ ?yr|per year|years?|yrs?|năm|nam|tháng)?\.?")


def _plain_number(tok):
    """Parse '25,000,000' / '25.000.000' / '1.234,56' / '1,5' / '8.5' (no unit)."""
    t = re.sub(r"\s+", "", tok)
    sign = -1 if t.startswith("-") else 1
    t = t.lstrip("+-")
    if "," in t and "." in t:
        dec = "," if t.rfind(",") > t.rfind(".") else "."
        t = t.replace("." if dec == "," else ",", "").replace(dec, ".")
    elif t.count(",") > 1 or t.count(".") > 1:
        t = t.replace(",", "").replace(".", "")
    elif "," in t:
        a, b = t.split(",")
        t = a + b if len(b) == 3 and 1 <= len(a) <= 3 and not a.startswith("0") else a + "." + b
    elif "." in t:
        a, b = t.split(".")
        t = a + b if len(b) == 3 and 1 <= len(a) <= 3 and not a.startswith("0") else t
    return sign * float(t)


def _clean_number(v):
    """Tolerant number reader for cell contents: numbers, '25,000,000 VND', '25.000.000đ', '25 triệu', '1.5 tỷ', '8%', '10 years'."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if not isinstance(v, str):
        return None
    s = v.replace("\u00a0", " ").strip().lower()
    m = _NUM_RE.fullmatch(s)
    if not m:
        return None
    try:
        x = _plain_number(m.group(1)) * _UNIT_X.get((m.group(2) or "").lower(), 1)
    except ValueError:
        return None
    return x if math.isfinite(x) else None


def sanitize_fields(d):
    """Keep only whitelisted, valid fields. Returns (clean, warnings)."""
    out, warns = {}, []
    for k, v in (d or {}).items():
        if k in ("liquidity", "risk"):
            w = CATEGORY_WORDS.get(_norm_label(v)) or CATEGORY_WORDS.get(_fold(v))
            if w:
                out[k] = w
            else:
                warns.append(f"'{k}' must be low, medium or high (got '{v}').")
        elif k in MONEY_KEYS or k in ("debt_rate", "years"):
            x = _clean_number(v)
            if x is None:
                warns.append(f"'{k}' is not a number ({str(v)[:30]!r}).")
            elif x < 0:
                warns.append(f"'{k}' cannot be negative; it was ignored.")
            elif k == "years" and not 1 <= x <= 30:
                warns.append("'years' must be between 1 and 30.")
            elif k == "debt_rate" and x > 100:
                warns.append("'debt_rate' must be between 0 and 100 (% per year).")
            elif x > MAX_VALUE:
                warns.append(f"'{k}' is unrealistically large.")
            else:
                out[k] = int(round(x)) if k == "years" else x
    return out, warns


def _norm_label(s):
    return re.sub(r"\s+", " ", str(s).strip().lower())


def template_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "FINOVA Input", [
        ("Field", "Value", "Example / allowed values"),
        ("Monthly income", 25000000, "VND per month, 0 or more"), ("Monthly expenses", 12000000, "VND per month, 0 or more"),
        ("Investments", 50000000, "VND, 0 or more"), ("Assets", 120000000, "VND, 0 or more"),
        ("Time horizon (years)", 10, "Whole number from 1 to 30"), ("Liquidity need", "medium", "low / medium / high"),
        ("Risk level", "medium", "low / medium / high"), ("Total current debt", 80000000, "VND, 0 or more"),
        ("Monthly debt payment", 3000000, "VND per month, 0 or more"),
        ("Debt interest rate", 0, "% per year, optional (0 to 100)")], widths=(28, 20, 40))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


FRIENDLY = {"income": "Monthly income", "expenses": "Monthly expenses", "investments": "Investments", "assets": "Assets", "years": "Time horizon (years)",
            "liquidity": "Liquidity need", "risk": "Risk level", "debt": "Total current debt", "debt_payment": "Monthly debt payment", "debt_rate": "Debt interest rate"}
GRID_ROWS, GRID_COLS = 250, 14


def _cell_value(c):
    v = c.value
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (c.number_format or "") and abs(v) <= 1.0:
        return f"{v * 100:g}%"                                   # 0.08 formatted as 8% means 8 per cent
    return v


def _grids_xlsx(raw):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        raise ValueError("The workbook is corrupted or password-protected and could not be read. Please save it again as .xlsx without a password.")
    grids = []
    for ws in wb.worksheets[:6]:
        if ws.sheet_state != "visible":
            continue
        try:
            rows = [[_cell_value(c) for c in row] for row in ws.iter_rows(min_row=1, max_row=GRID_ROWS, max_col=GRID_COLS)]
        except Exception:
            continue
        rows = [[None if (x is None or (isinstance(x, str) and not x.strip())) else x for x in r] for r in rows]
        if any(any(x is not None for x in r) for r in rows):
            grids.append((ws.title, rows))
    return grids


def _grids_csv(raw):
    import csv
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("The text file could not be decoded. Please save it as UTF-8 CSV.")
    try:
        dialect = csv.Sniffer().sniff(text[:3000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[(x.strip() or None) for x in r][:GRID_COLS] for r in csv.reader(io.StringIO(text), dialect)][:GRID_ROWS]
    return [("CSV", rows)] if any(any(x is not None for x in r) for r in rows) else []


def parse_xlsx(raw, filename=""):
    """Read a financial profile from an uploaded workbook (.xlsx/.xlsm) or CSV, whatever its layout:
      * label in one cell and value to its right (the FINOVA template, or any 'Label | Value' list)
      * a header row of labels with the values in the row below (a horizontal table)
      * labels in English or Vietnamese, with or without accents, units or notes in brackets
      * values such as 25000000, '25,000,000 VND', '25.000.000đ', '25 triệu', '1.5 tỷ', 8%, '10 years'
    Raises ValueError with a friendly message when the file cannot be used."""
    if not raw:
        raise ValueError("The file is empty.")
    if len(raw) > MAX_XLSX:
        raise ValueError("The file is larger than 4 MB.")
    if raw[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError("This is the old Excel format (.xls) or an encrypted file. In Excel use File > Save As > Excel Workbook (.xlsx), then upload it again.")
    if raw[:2] == b"PK":
        grids, kind = _grids_xlsx(raw), "xlsx"
    elif filename.lower().endswith((".csv", ".txt")):
        grids, kind = _grids_csv(raw), "csv"
    else:
        raise ValueError("This is not a valid .xlsx workbook (wrong file signature). Use the FINOVA Excel template, or a .xlsx / .csv file.")
    if not grids:
        raise ValueError("The file has no data. Fill in the Value column of the FINOVA template and upload it again.")
    data, warnings, seen, unrecognized, rows_for_ai = {}, [], set(), [], []
    for sheet, G in grids:
        for r, row in enumerate(G):
            labelled = [(c, field_for_label(v)) for c, v in enumerate(row) if isinstance(v, str) and field_for_label(v)]
            cols = [c for c, _ in labelled]
            header_row = len({k for _, k in labelled}) >= 3 or (len(cols) >= 2 and any(c + 1 in cols for c in cols))     # a row of labels: values are below
            for c, key in labelled:
                label = str(row[c]).strip()[:60]
                val = None
                if ":" in label and _fold(label.split(":")[0]) in FIELD_MAP:                    # "Monthly income: 25,000,000"
                    val = label.split(":", 1)[1].strip() or None
                if val is None and header_row and r + 1 < len(G):                                   # horizontal table: value is below the label
                    val = G[r + 1][c]
                elif val is None and c + 1 < len(row):                                              # label | value: value is the next cell
                    val = row[c + 1]
                if key in seen:
                    warnings.append(f"'{label}' appears more than once; the first value was used.")
                    continue
                seen.add(key)
                if val is None:
                    warnings.append(f"'{label}' has no value, so it was left empty.")
                    continue
                if isinstance(val, str) and val.startswith("#"):
                    warnings.append(f"'{label}' contains a spreadsheet error ({val[:12]}). Fix the cell and upload again.")
                    continue
                clean, w = sanitize_fields({key: val})
                warnings += [f"'{label}': {x}" for x in w]
                data.update(clean)
            if not labelled and not header_row and len(row) > 1 and isinstance(row[0], str) and row[1] is not None:
                lab = row[0].strip()[:60]
                if _fold(lab) and _fold(lab) not in HEADER_WORDS:
                    unrecognized.append({"label": lab, "value": row[1] if isinstance(row[1], (int, float)) and not isinstance(row[1], bool) else str(row[1])[:60]})
            if len(rows_for_ai) < 60 and len(row) > 1 and isinstance(row[0], str) and row[1] is not None:
                rows_for_ai.append((row[0].strip()[:60], row[1] if isinstance(row[1], (int, float)) else str(row[1])[:60]))
    if not data and not unrecognized:
        raise ValueError("No financial labels were found. Use labels such as 'Monthly income' and 'Monthly expenses' (the FINOVA template has them all).")
    unrec = [u for u in unrecognized if FIELD_MAP.get(_fold(u["label"])) is None][:20]
    missing = [k for k in ALIASES if k not in data and k != "debt_rate"]
    return {"data": data, "warnings": warnings, "missing": missing, "missing_labels": [FRIENDLY[k] for k in missing], "unrecognized": unrec, "rows": rows_for_ai, "format": kind}


# ------------------------------------------------------------- grounding
def _numbers_in(obj, out):
    if isinstance(obj, bool):
        return
    if isinstance(obj, (int, float)):
        out.append(float(obj))
    elif isinstance(obj, str):                       # figures quoted inside Python-written explanations are engine figures too
        for t in re.findall(r"\d[\d,]*\.?\d*", obj[:2000]):
            try:
                out.append(float(t.replace(",", "")))
            except ValueError:
                pass
    elif isinstance(obj, dict):
        for v in obj.values():
            _numbers_in(v, out)
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            _numbers_in(v, out)


def collect_numbers(obj):
    out = []
    _numbers_in(obj, out)
    return out


def grounded_rate(text, outputs, money_only=False):
    """(grounded, checked): money (comma-formatted) and percentage figures in `text` that exist in
    the engine outputs. Used for the Groundedness evaluation and to reject hallucinated LLM text."""
    vals = collect_numbers(outputs)
    checked = ok = 0
    for t in re.findall(r"\d[\d,]*\.?\d*%?", str(text)):
        pct_ = t.endswith("%")
        try:
            x = float(t.rstrip("%").replace(",", ""))
        except ValueError:
            continue
        if not pct_ and "," not in t:
            continue
        if pct_ and money_only:
            continue
        checked += 1
        if pct_:
            hit = any(abs(abs(v) * 100 - x) <= 0.06 or abs(abs(v) - x) <= 0.06 for v in vals)
        else:
            hit = any(abs(abs(v) - x) <= max(1.0, 0.002 * abs(v)) for v in vals)
        ok += hit
    return ok, checked
